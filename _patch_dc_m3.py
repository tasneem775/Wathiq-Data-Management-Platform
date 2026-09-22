"""Patch DC.M.3.docx — fix executive summary text that contradicts its own KPI table.

Table 10 in the document shows: DC-KPI-01 محقق, DC-KPI-02 محقق,
DC-KPI-03 قريب من التحقيق, DC-KPI-04 قريب من التحقيق (2 محقق / 2 قريب من التحقيق).
The section-2 narrative wrongly claimed "3 محقق / 1 قريب من التحقيق".
"""
import sys; sys.stdout.reconfigure(encoding="utf-8")
from docx import Document

PATH = (r"c:\Users\extra\Downloads\NDI-Sentinel"
        r"\evidence_repository\data_classification\MQ1\DC.M.3.docx")

doc = Document(PATH)

OLD = ("حققت الهيئة نتائج إيجابية عبر جميع المؤشرات الأربعة المرصودة، إذ بلغت ثلاثة "
       "مؤشرات منها مستوى (محقق) ومؤشر واحد وصل إلى مستوى (قريب من التحقيق).")
NEW = ("حققت الهيئة نتائج إيجابية عبر جميع المؤشرات الأربعة المرصودة، إذ بلغ مؤشرا "
       "مجموعات البيانات المصنفة وتعيين ملاك البيانات مستوى (محقق)، فيما وصل مؤشرا "
       "إكمال التوعية والالتزام بالإجراءات إلى مستوى (قريب من التحقيق).")

changed = False
for para in doc.paragraphs:
    if "بلغت ثلاثة مؤشرات منها مستوى" in para.text:
        full = para.text.strip()
        with open("_debug_dc_m3.txt", "w", encoding="utf-8") as f:
            f.write("FOUND:\n" + full + "\n\nEXPECTED:\n" + OLD + "\n")
        runs = [r for r in para.runs if r.text.strip()]
        if runs:
            runs[0].text = NEW
            for r in runs[1:]:
                r.text = ""
            changed = True

if not changed:
    print("ERROR: target paragraph not found")
else:
    doc.save(PATH)
    print("Patched OK")

doc2 = Document(PATH)
for para in doc2.paragraphs:
    if "مؤشرا مجموعات البيانات" in para.text or "قريب من التحقيق" in para.text:
        print("VERIFY:", para.text)
