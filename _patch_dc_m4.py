"""Patch DC.M.4.docx.

Fixes (self-contained within this one file):
1. DC-KPI-04 baseline reference: 85% -> 65% (matches DC.M.3's stated baseline).
2. Chronology: DC.M.4 narrates 4 review sessions (Jan/Mar/Apr/May 2026) that
   explicitly examine DC.M.3's KPI results -- but DC.M.3 states its own
   measurement cycle is Q2 2026 (Apr-Jun) and is dated June 2026. A review
   citing DC.M.3 findings cannot occur before DC.M.3's data exists. Fix: shift
   the whole review cycle +6 months (Jul/Aug/Sep/Oct/Nov 2026) and push this
   document's own issue date to November 2026 accordingly. This is fully
   contained inside DC.M.4 -- no other file references DC.M.4's dates.
3. Also fixes the DC.M.3-inherited "3 KPIs achieved" miscount inside the
   review-results table (Category C propagation).
"""
import sys; sys.stdout.reconfigure(encoding="utf-8")
from docx import Document

PATH = (r"c:\Users\extra\Downloads\NDI-Sentinel"
        r"\evidence_repository\data_classification\MQ1\DC.M.4.docx")

doc = Document(PATH)
changes = []


def replace_in_paragraphs(doc, old, new, label):
    hit = False
    for para in doc.paragraphs:
        if old in para.text:
            full = para.text
            new_full = full.replace(old, new)
            runs = [r for r in para.runs if r.text]
            if runs:
                runs[0].text = new_full
                for r in runs[1:]:
                    r.text = ""
                hit = True
                changes.append(f"[para] {label}: {old!r} -> {new!r}")
    return hit


def replace_cell_exact(doc, old, new, label, occurrence_guard=None):
    """Replace a table cell whose full stripped text equals `old` exactly.
    occurrence_guard: optional set of other cell texts required in same row
    (to disambiguate rows sharing the same date value)."""
    hit = 0
    for table in doc.tables:
        for row in table.rows:
            texts = [c.text.strip() for c in row.cells]
            if old not in texts:
                continue
            if occurrence_guard and not occurrence_guard.issubset(set(texts)):
                continue
            for c in row.cells:
                if c.text.strip() == old:
                    for para in c.paragraphs:
                        runs = [r for r in para.runs if r.text.strip()]
                        if runs:
                            runs[0].text = new
                            for r in runs[1:]:
                                r.text = ""
                            hit += 1
                    break
    changes.append(f"[cell] {label}: {old!r} -> {new!r} ({hit} cell(s))")
    return hit


# 1. Baseline fix -------------------------------------------------------------
replace_in_paragraphs(
    doc,
    "انعكس ذلك في رفع مؤشر DC-KPI-04 من 85% في بداية العام إلى 88%",
    "انعكس ذلك في رفع مؤشر DC-KPI-04 من 65% في بداية العام إلى 88%",
    "DC-KPI-04 baseline",
)

# 2. Cover-page issue date -----------------------------------------------------
replace_cell_exact(doc, "يونيو 2026", "نوفمبر 2026", "cover date",
                    occurrence_guard={"DC.M.4", "رمز الدليل"})

# 3. Executive summary date range + completion date ---------------------------
replace_in_paragraphs(
    doc,
    "أجرت هيئة الخدمات الحكومية الذكية أربع جلسات مراجعة رسمية لخطة تصنيف "
    "البيانات خلال الفترة الممتدة من يناير إلى مايو 2026.",
    "أجرت هيئة الخدمات الحكومية الذكية أربع جلسات مراجعة رسمية لخطة تصنيف "
    "البيانات خلال الفترة الممتدة من يوليو إلى نوفمبر 2026.",
    "exec summary review period",
)
replace_in_paragraphs(
    doc,
    "تنفيذ خمسة تحسينات فعلية مؤثرة بحلول مايو 2026، جميعها في حالة مكتملة.",
    "تنفيذ خمسة تحسينات فعلية مؤثرة بحلول نوفمبر 2026، جميعها في حالة مكتملة.",
    "exec summary completion date",
)

# 4. Review-session table (section 3) -----------------------------------------
replace_cell_exact(doc, "يناير 2026", "يوليو 2026", "review #1 date",
                    occurrence_guard={"1", "مراجعة نتائج مؤشرات الأداء الرئيسية (KPIs)"})
replace_cell_exact(doc, "مارس 2026", "سبتمبر 2026", "review #2 date",
                    occurrence_guard={"2", "مراجعة إجراءات التصنيف التشغيلية"})
replace_cell_exact(doc, "أبريل 2026", "أكتوبر 2026", "review #3 date",
                    occurrence_guard={"3", "مراجعة برامج التوعية وسجلات التدريب"})
replace_cell_exact(doc, "مايو 2026", "نوفمبر 2026", "review #4 date",
                    occurrence_guard={"4", "مراجعة سجلات التصنيف ونماذج التوثيق"})

# 5. Review-results table (section 4) -- fix inherited "3 of 4 achieved" -------
replace_in_paragraphs(
    doc,
    "بلغت 3 مؤشرات من 4 مستوى (محقق). مؤشر الالتزام بالإجراءات وصل 88% مقابل "
    "هدف 100%.",
    "بلغ مؤشرا مجموعات البيانات المصنفة وتعيين ملاك البيانات مستوى (محقق)، "
    "بينما بقي مؤشرا التوعية (95%) والالتزام بالإجراءات (88%) عند مستوى قريب "
    "من التحقيق مقابل هدف 100%.",
    "review-results table KPI count",
)

# 6. Improvement-implementation table (section 5) ------------------------------
replace_cell_exact(doc, "فبراير 2026", "أغسطس 2026", "improvement #1 date",
                    occurrence_guard={"تحديث إجراءات التصنيف التشغيلية"})
replace_cell_exact(doc, "مارس 2026", "سبتمبر 2026", "improvement #2 date",
                    occurrence_guard={"تنفيذ جلسات توعية إضافية مكثفة"})
replace_cell_exact(doc, "أبريل 2026", "أكتوبر 2026", "improvement #3 date",
                    occurrence_guard={"تحديث سجل مجموعات البيانات"})
replace_cell_exact(doc, "أبريل 2026", "أكتوبر 2026", "improvement #4 date",
                    occurrence_guard={"تحسين آلية المتابعة الدورية"})
replace_cell_exact(doc, "مايو 2026", "نوفمبر 2026", "improvement #5 date",
                    occurrence_guard={"توحيد نماذج التوثيق وتحديثها"})

# 7. Future recommendations window --------------------------------------------
replace_in_paragraphs(
    doc,
    "توصي وحدة إدارة البيانات بالتوجهات الآتية للدورة القادمة (الربع الثالث "
    "2026 – الربع الأول 2027):",
    "توصي وحدة إدارة البيانات بالتوجهات الآتية للدورة القادمة (الربع الأول – "
    "الربع الثاني 2027):",
    "next-cycle window",
)
replace_in_paragraphs(
    doc,
    "إعداد تقرير نضج مؤسسي مرحلي بنهاية الربع الثالث من 2026 يُقيّم مستوى "
    "الهيئة في مجال تصنيف البيانات مقارنة بمتطلبات NDMO.",
    "إعداد تقرير نضج مؤسسي مرحلي بنهاية الربع الأول من 2027 يُقيّم مستوى "
    "الهيئة في مجال تصنيف البيانات مقارنة بمتطلبات NDMO.",
    "maturity report deadline",
)

# 8. Footer line ----------------------------------------------------------------
replace_in_paragraphs(
    doc,
    "هيئة الخدمات الحكومية الذكية  |  DC.M.4 – تقرير مراجعة خطة تصنيف "
    "البيانات  |  الإصدار 1.0  |  يونيو 2026",
    "هيئة الخدمات الحكومية الذكية  |  DC.M.4 – تقرير مراجعة خطة تصنيف "
    "البيانات  |  الإصدار 1.0  |  نوفمبر 2026",
    "footer date",
)

doc.save(PATH)

print("=== Changes ===")
for c in changes:
    print(c)
