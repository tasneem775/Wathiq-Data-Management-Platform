"""Patch DC.M.5.docx — update 2 classification cells + summary table."""
import sys; sys.stdout.reconfigure(encoding="utf-8")
from docx import Document

PATH = (r"c:\Users\extra\Downloads\NDI-Sentinel"
        r"\evidence_repository\data_classification\MQ2\DC.M.5.docx")

doc = Document(PATH)

changes = []

# ── helper: update first run text in a cell (preserves formatting) ──────────
def set_cell_text(cell, new_text):
    """Replace text in a cell while keeping run formatting intact."""
    for para in cell.paragraphs:
        full = para.text.strip()
        if full:
            # Write into first non-empty run, clear the rest
            runs = [r for r in para.runs if r.text.strip()]
            if runs:
                runs[0].text = new_text
                for r in runs[1:]: r.text = ""
            return True
    return False

# ══════════════════════════════════════════════════════════════════════════════
# Change 1 & 2: classification level cells in the main dataset table
#
# With bidiVisual the XML cell order is reversed from visual order.
# The table was built with mk_tbl so visual RTL cols include التصنيف on the left.
# We scan every row for a cell whose text matches the dataset name,
# then locate the classification cell in the same row by its current value.
# ══════════════════════════════════════════════════════════════════════════════
TARGETS = {
    "بيانات المستخدمين":    {"old": "مقيّد", "new": "سرّي"},
    "بيانات الأصول التقنية": {"old": "مقيّد", "new": "سرّي"},
}

for table in doc.tables:
    for row in table.rows:
        cell_texts = [c.text.strip() for c in row.cells]
        for name, chg in TARGETS.items():
            if name in cell_texts and chg["old"] in cell_texts:
                # find the cell that holds the old classification value
                for c in row.cells:
                    if c.text.strip() == chg["old"]:
                        set_cell_text(c, chg["new"])
                        changes.append(
                            f"  ✓  [{name}]  {chg['old']}  →  {chg['new']}")
                        break   # only the first matching cell in this row

# ══════════════════════════════════════════════════════════════════════════════
# Change 3: summary table values  (مقيّد 7→5, سرّي 7→9)
# We scan for cells containing the exact old counts and patch them.
# ══════════════════════════════════════════════════════════════════════════════
SUMMARY_PATCHES = {
    # (label_text, old_count) → new_count
    # We match on both the label cell and the value cell being in the same row.
    "مقيّد": ("7", "5"),
    "سرّي":  ("7", "9"),
}

for table in doc.tables:
    for row in table.rows:
        cell_texts = [c.text.strip() for c in row.cells]
        for label, (old_val, new_val) in SUMMARY_PATCHES.items():
            if label in cell_texts and old_val in cell_texts:
                for c in row.cells:
                    if c.text.strip() == old_val:
                        set_cell_text(c, new_val)
                        changes.append(
                            f"  ✓  ملخص [{label}]  {old_val}  →  {new_val}")
                        break

doc.save(PATH)

# ── Verification ──────────────────────────────────────────────────────────────
print("=== Changes applied ===")
for ch in changes:
    print(ch)

print("\n=== Verification ===")
doc2 = Document(PATH)
for table in doc2.tables:
    for row in table.rows:
        texts = [c.text.strip() for c in row.cells]
        # Show rows with our target names or summary counts
        if any(t in texts for t in
               ["بيانات المستخدمين","بيانات الأصول التقنية",
                "مقيّد","سرّي","عام","سرّي للغاية","الإجمالي"]):
            print("  ROW:", texts)
