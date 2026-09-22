# -*- coding: utf-8 -*-
"""Builds the supporting Excel calculation sheet for DC.C.4.1-E02 (KPI-DC-02) — 7-section
compliance-audit template. All per-asset rows are copied verbatim from the real
DC.C.3.2.docx and DC.C.5.1.docx (Table 6). No asset name, level, or count is invented."""
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT_PATH = "06_DC_C4_1_KPI_Evidence/DC.C.4.1-E02_KPI-DC-02_Classification_Distribution_KPI_Calculation_Record.xlsx"

NAVY = "1F2A44"
GOLD = "AB8654"
LIGHT_GREY = "F2F2F2"
WHITE = "FFFFFF"
GREEN = "1F6F43"
AMBER = "8A6D00"
AR_FONT = "Arial"

ASSETS = [
    (1, "بيانات الهوية والتحقق الإلكتروني", "سرّي للغاية", "سرّي للغاية"),
    (2, "بيانات الأحداث والتنبيهات الأمنية", "سرّي للغاية", "سرّي للغاية"),
    (3, "بيانات الصلاحيات وإدارة الهوية", "سرّي للغاية", "سرّي للغاية"),
    (4, "بيانات التكامل مع الجهات الحكومية", "سرّي للغاية", "سرّي للغاية"),
    (5, "بيانات الرواتب والمزايا الوظيفية", "سرّي", "سرّي"),
    (6, "بيانات المدفوعات والمعاملات المالية", "سرّي", "سرّي"),
    (7, "بيانات الموظفين والملفات الوظيفية", "سرّي", "سرّي"),
    (8, "بيانات المستخدمين والمستفيدين", "سرّي", "سرّي"),
    (9, "بيانات طلبات ومعاملات المستفيدين", "مقيّد", "مقيّد"),
    (10, "بيانات الخدمات الحكومية الرقمية", "مقيّد", "عام"),
    (11, "بيانات الموردين والمتعاقدين", "مقيّد", "مقيّد"),
    (12, "بيانات الأصول والبنية التحتية التقنية", "سرّي", "سرّي"),
    (13, "بيانات الأهداف والخطط الاستراتيجية", "سرّي", "سرّي"),
    (14, "بيانات التدريب والتطوير الوظيفي", "مقيّد", "مقيّد"),
    (15, "بيانات استطلاعات رضا المستفيدين", "عام", "عام"),
    (16, "بيانات الحوادث التقنية وطلبات الدعم", "مقيّد", "مقيّد"),
    (17, "بيانات المحتوى الرقمي والموقع الإلكتروني", "عام", "عام"),
    (18, "بيانات التقارير والإحصاءات المنشورة", "عام", "عام"),
    (19, "سجلات العقود والاتفاقيات", "سرّي", "سرّي"),
    (20, "سجلات التدقيق الداخلي والرقابة", "سرّي", "سرّي"),
    (21, "سجلات المراسلات والخطابات الرسمية", "مقيّد", "مقيّد"),
    (22, "سجلات الاجتماعات واللجان الحوكمية", "مقيّد", "مقيّد"),
    (23, "سجلات التقارير الإدارية الدورية", "مقيّد", "عام"),
    (24, "سجلات فعاليات الهيئة وأحداثها", "عام", "عام"),
    (25, "سجلات الشكاوى والمقترحات", "مقيّد", "مقيّد"),
    (26, "سجلات قرارات وتوجيهات الإدارة العليا", "سرّي", "سرّي"),
    (27, "سجلات التراخيص والتصاريح القانونية", "مقيّد", "عام"),
    (28, "سجلات النسخ الاحتياطي واسترجاع البيانات", "سرّي", "سرّي"),
    (29, "وثائق السياسات والإجراءات الداخلية", "مقيّد", "مقيّد"),
    (30, "سجلات التغييرات والتحديثات التقنية", "مقيّد", "مقيّد"),
]
TOTAL_ASSETS = len(ASSETS)
dist = Counter(level_after for _, _, _, level_after in ASSETS)
DIST_ORDER = ["سرّي للغاية", "سرّي", "مقيّد", "عام"]
CLASS_COLOR = {"سرّي للغاية": NAVY, "سرّي": AMBER, "مقيّد": AMBER, "عام": GREEN}


def pct(n):
    return round(n / TOTAL_ASSETS * 1000) / 10


wb = Workbook()
ws = wb.active
ws.title = "احتساب KPI-DC-02"
ws.sheet_view.showGridLines = False
ws.sheet_view.rightToLeft = True

thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

N_COLS = 6
r = 1


def merged_text(row, text, size, bold, color, fill=None, height=None, italic=False, span=N_COLS):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name=AR_FONT, size=size, bold=bold, italic=italic, color=color)
    cell.alignment = Alignment(horizontal="right", vertical="center", indent=1, wrap_text=True)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if height:
        ws.row_dimensions[row].height = height
    return row + 1


def section_heading(row, text, span=N_COLS):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name=AR_FONT, size=12.5, bold=True, color=NAVY)
    cell.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    cell.border = Border(bottom=Side(style="medium", color=GOLD))
    ws.row_dimensions[row].height = 22
    return row + 1


def spacer(row, h=10):
    ws.row_dimensions[row].height = h
    return row + 1


def meta_row(row, label, value):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    c1 = ws.cell(row=row, column=1, value=label)
    c1.font = Font(name=AR_FONT, size=10, bold=True, color="595959")
    c1.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=N_COLS)
    c2 = ws.cell(row=row, column=3, value=value)
    c2.font = Font(name=AR_FONT, size=10, color="262626")
    c2.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    return row + 1


def bullet_row(row, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=N_COLS)
    cell = ws.cell(row=row, column=1, value="•  " + text)
    cell.font = Font(name=AR_FONT, size=9.5, color="262626")
    cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True, indent=1)
    ws.row_dimensions[row].height = 30
    return row + 1


r = merged_text(r, "سجل احتساب توزيع الأصول حسب مستوى التصنيف (KPI-DC-02)", 18, True, WHITE, fill=NAVY, height=32)
r = merged_text(r, "منصة ذكية لإدارة وقياس الامتثال لمتطلبات المكتب الوطني لإدارة البيانات (NDMO)", 10.5, True, WHITE, fill=NAVY, height=20)
r = spacer(r, 6)
r = merged_text(r, "سجل احتساب توزيع الأصول حسب مستوى التصنيف (KPI-DC-02)", 14, True, NAVY, height=24)
r = merged_text(r, "دليل داعم — DC.C.4.1-E02 — Calculation Record", 11, True, GOLD, italic=True, height=18)
r = spacer(r, 10)

r = section_heading(r, "1. بيانات الوثيقة")
r = spacer(r, 4)
r = meta_row(r, "الجهة", "هيئة الخدمات الحكومية الذكية (SGSA)")
r = meta_row(r, "المنصة", "")
r = meta_row(r, "الإطار المرجعي", "إطار حوكمة البيانات الوطني (NDMO)")
r = meta_row(r, "رمز الدليل", "DC.C.4.1-E02")
r = meta_row(r, "نوع الوثيقة", "Calculation Record")
r = meta_row(r, "المتطلب المرتبط", "DC.C.4.1 — تقرير مراقبة عمليات تصنيف البيانات عبر مؤشرات الأداء الرئيسية")
r = meta_row(r, "المؤشر المرتبط", "KPI-DC-02 — نسبة الأصول المصنفة بكل مستوى تصنيف")
r = meta_row(r, "الإصدار", "1.0")
r = meta_row(r, "التاريخ", "يونيو 2026")
r = meta_row(r, "حالة الوثيقة", "جاهزة للاعتماد")
r = spacer(r, 14)

r = section_heading(r, "2. نطاق القياس")
r = spacer(r, 4)
r = meta_row(r, "فترة القياس", "الربع الثاني 2026 (أبريل – يونيو 2026)")
r = meta_row(r, "نطاق القياس", "الأصول الثلاثون بمستوى تصنيفها النهائي المعتمد أثناء المراجعة (DC.C.5.1)")
r = spacer(r, 14)

r = section_heading(r, "3. مصدر البيانات")
r = spacer(r, 4)
r = bullet_row(r, "DC.C.3.2 — تقرير تقييم الأثر (مستوى التصنيف الممنوح أصلاً)")
r = bullet_row(r, "DC.C.5.1 — سجل البيانات (مستوى التصنيف المعتمد أثناء المراجعة)")
r = spacer(r, 14)

r = section_heading(r, "4. جدول البيانات المستخدمة (مصدرها DC.C.3.2 وDC.C.5.1)")
r = spacer(r, 4)

headers = ["م", "اسم الأصل", "التصنيف الممنوح", "التصنيف المعتمد", "معادلة الحساب", "النسبة النهائية"]
header_row = r
for c, h in enumerate(headers, start=1):
    cell = ws.cell(row=header_row, column=c, value=h)
    cell.font = Font(name=AR_FONT, size=10, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border
ws.row_dimensions[header_row].height = 34
r += 1

for i, (n, name, granted, final) in enumerate(ASSETS):
    row_fill = LIGHT_GREY if i % 2 == 0 else WHITE
    formula_txt = f"{dist[final]} ÷ {TOTAL_ASSETS} × 100"
    result_txt = f"{pct(dist[final])}%"
    row_vals = [n, name, granted, final, formula_txt, result_txt]
    for c, v in enumerate(row_vals, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.border = border
        cell.fill = PatternFill("solid", fgColor=row_fill)
        center_cols = (1, 3, 4, 6)
        cell.alignment = Alignment(
            horizontal="center" if c in center_cols else "right",
            vertical="center", wrap_text=True,
            indent=0 if c in center_cols else 1,
        )
        if c == 1:
            cell.font = Font(name=AR_FONT, size=10, bold=True, color=NAVY)
        elif c in (3, 4):
            cell.font = Font(name=AR_FONT, size=10, bold=True, color=CLASS_COLOR.get(v, "262626"))
        else:
            cell.font = Font(name=AR_FONT, size=9.5, color="262626")
    ws.row_dimensions[r].height = 26
    r += 1

r = spacer(r, 14)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=N_COLS)
note1 = ws.cell(
    row=r, column=1,
    value=(
        "ملاحظة تدقيق: الأصول رقم 10 و23 و27 هي الأصول التي أُعيد تصنيفها من «مقيّد» إلى "
        "«عام» بموجب مراجعة DC.C.3.3. لم تُستخدم بيانات DC.M.7 لأنها تعكس التصنيف قبل هذه "
        "المراجعة."
    ),
)
note1.font = Font(name=AR_FONT, size=8, italic=True, color="7F7F7F")
note1.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True, indent=1)
ws.row_dimensions[r].height = 34
r += 1

r = spacer(r, 14)
r = section_heading(r, "5. تطبيق المعادلة")
r = spacer(r, 4)

summary_headers = ["مستوى التصنيف", "عدد الأصول", "تطبيق المعادلة", "النسبة النهائية"]
sh_row = r
col_spans2 = [(1, 1), (2, 2), (3, 4), (5, 6)]
for (sc, ec), h in zip(col_spans2, summary_headers):
    ws.merge_cells(start_row=sh_row, start_column=sc, end_row=sh_row, end_column=ec)
    cell = ws.cell(row=sh_row, column=sc, value=h)
    cell.font = Font(name=AR_FONT, size=10, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    for cc in range(sc, ec + 1):
        ws.cell(row=sh_row, column=cc).border = border
ws.row_dimensions[sh_row].height = 20
r += 1

for level in DIST_ORDER + ["الإجمالي"]:
    n = TOTAL_ASSETS if level == "الإجمالي" else dist.get(level, 0)
    formula = f"{n} ÷ {TOTAL_ASSETS} × 100"
    result = "100%" if level == "الإجمالي" else f"{pct(n)}%"
    vals = [level, str(n), formula, result]
    for (sc, ec), v in zip(col_spans2, vals):
        ws.merge_cells(start_row=r, start_column=sc, end_row=r, end_column=ec)
        cell = ws.cell(row=r, column=sc, value=v)
        cell.font = Font(name=AR_FONT, size=10, bold=(sc in (1, 5)), color=CLASS_COLOR.get(level, "262626") if sc == 1 else "262626")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        for cc in range(sc, ec + 1):
            ws.cell(row=r, column=cc).border = border
    ws.row_dimensions[r].height = 24
    r += 1

r = spacer(r, 14)
r = section_heading(r, "6. النتيجة والتحقق مقابل DC.C.4.1")
r = spacer(r, 4)
r = bullet_row(r, "النسب الناتجة مطابقة تماماً للقيم الواردة في بطاقة KPI-DC-02 ضمن تقرير DC.C.4.1 (القسم 8 وبند 9.2).")
r = bullet_row(r, "تم التحقق من أن مجموع الأصول عبر المستويات الأربعة (4+10+9+7) يساوي 30.")
r = spacer(r, 14)

r = section_heading(r, "7. الاعتماد والتوقيع")
r = spacer(r, 4)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=N_COLS)
review_cell = ws.cell(
    row=r, column=1,
    value="تمت مراجعة هذا السجل من مكتب إدارة البيانات، وتم التحقق من مطابقة البيانات الواردة فيه لمحتوى DC.C.3.2 وDC.C.5.1 دون أي إضافة أو حذف أو تعديل.",
)
review_cell.font = Font(name=AR_FONT, size=10, color="262626")
review_cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True, indent=1)
ws.row_dimensions[r].height = 26
r += 1
r = spacer(r, 6)

approval_headers = ["صاحب الصلاحية", "المسمى الوظيفي", "التوقيع", "التاريخ"]
col_spans = [(1, 1), (2, 2), (3, 4), (5, 6)]
approval_header_row = r
for (sc, ec), h in zip(col_spans, approval_headers):
    ws.merge_cells(start_row=approval_header_row, start_column=sc, end_row=approval_header_row, end_column=ec)
    cell = ws.cell(row=approval_header_row, column=sc, value=h)
    cell.font = Font(name=AR_FONT, size=10, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    for cc in range(sc, ec + 1):
        ws.cell(row=approval_header_row, column=cc).border = border
ws.row_dimensions[approval_header_row].height = 20
r += 1

approval_values = ["____________", "مدير مكتب إدارة البيانات", "____________", "____________"]
for (sc, ec), v in zip(col_spans, approval_values):
    ws.merge_cells(start_row=r, start_column=sc, end_row=r, end_column=ec)
    cell = ws.cell(row=r, column=sc, value=v)
    cell.font = Font(name=AR_FONT, size=9.5, color="262626")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    for cc in range(sc, ec + 1):
        ws.cell(row=r, column=cc).border = border
ws.row_dimensions[r].height = 26

widths = [6, 40, 16, 16, 24, 14]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

wb.save(OUT_PATH)
print("Saved:", OUT_PATH)
print("Distribution:", dict(dist))
