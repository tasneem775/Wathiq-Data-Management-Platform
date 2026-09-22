# -*- coding: utf-8 -*-
"""Builds a standalone compliance-audit Calculation Record for DC.C.4.1 (KPI-DC-03).

Purpose: document the data source and calculation method behind the "75%" value of
KPI-DC-03 (نسبة البيانات منخفضة الأثر المصنفة "مقيّد") as reported in DC.C.4.1.docx.

The 12-asset scope and per-asset final decision are copied verbatim from the real
DC.C.3.3.docx (Table 14 — القرار النهائي لكل أصل)، cross-referenced with the original
serial numbers from DC.C.3.2.docx (Table 11 — الأصول المصنَّفة مقيّداً أصلاً). No asset
name, decision, or count is invented.

Audit template (7 sections, per compliance-auditor brief):
1. بيانات الوثيقة  2. نطاق القياس  3. مصدر البيانات  4. جدول البيانات المستخدمة
5. تطبيق المعادلة  6. النتيجة والتحقق مقابل DC.C.4.1  7. الاعتماد والتوقيع
"""
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_PATH = "06_DC_C4_1_KPI_Evidence/DC.C.4.1-E03_KPI-DC-03_Low_Impact_Data_Review_KPI_Calculation_Record.docx"

NAVY = RGBColor(0x1F, 0x2A, 0x44)
GOLD = RGBColor(0xAB, 0x8D, 0x54)
GREY = RGBColor(0x59, 0x59, 0x59)
GREEN = RGBColor(0x1F, 0x6F, 0x43)
AMBER = RGBColor(0x8A, 0x6D, 0x00)
AR_FONT = "Arial"


def rtl(p):
    b = OxmlElement("w:bidi")
    b.set(qn("w:val"), "1")
    p._p.get_or_add_pPr().append(b)


def set_rtl_run(run):
    pr = run._element.get_or_add_rPr()
    pr.get_or_add_rFonts().set(qn("w:cs"), AR_FONT)
    e = OxmlElement("w:rtl")
    e.set(qn("w:val"), "1")
    pr.append(e)


def set_table_rtl(table):
    tbl_pr = table._tbl.tblPr
    bv = OxmlElement("w:bidiVisual")
    bv.set(qn("w:val"), "1")
    tbl_pr.append(bv)


def set_cell_background(cell, color_hex):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color_hex)
    cell._tc.get_or_add_tcPr().append(shd)


def add_meta_row(table, label, value):
    row = table.add_row()
    p1 = row.cells[1].paragraphs[0]
    rtl(p1)
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p1.add_run(label)
    r1.font.bold = True
    r1.font.size = Pt(9)
    r1.font.color.rgb = GREY
    set_rtl_run(r1)

    p2 = row.cells[0].paragraphs[0]
    rtl(p2)
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = p2.add_run(value)
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
    set_rtl_run(r2)


def add_paragraph_ar(doc, text, size=9.5, bold=False, italic=False, color=RGBColor(0x26, 0x26, 0x26)):
    p = doc.add_paragraph()
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    set_rtl_run(r)
    return p


def add_heading_ar(doc, text, size=13):
    p = doc.add_paragraph()
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = NAVY
    set_rtl_run(r)
    return p


def add_bullet_ar(doc, text, size=9.5):
    p = doc.add_paragraph()
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("•  " + text)
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
    set_rtl_run(r)
    return p


DECISION_COLOR = {"الإبقاء كمقيّد": AMBER, "إعادة التصنيف إلى عام": GREEN}


def make_table(doc, headers, rows, col_widths, status_col=None, font_size=8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_rtl(table)
    for i, w in enumerate(col_widths):
        table.columns[i].width = w

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        set_cell_background(hdr_cells[i], "1F2A44")
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = hdr_cells[i].paragraphs[0]
        rtl(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_rtl_run(r)

    for idx, row_vals in enumerate(rows):
        row = table.add_row()
        for c, v in enumerate(row_vals):
            cell = row.cells[c]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if idx % 2 == 0:
                set_cell_background(cell, "F2F2F2")
            p = cell.paragraphs[0]
            rtl(p)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c in (status_col or []) else WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(str(v))
            run.font.size = Pt(font_size)
            set_rtl_run(run)
            if c == 0:
                run.font.bold = True
                run.font.color.rgb = NAVY
            elif status_col and c in status_col:
                run.font.bold = True
                run.font.color.rgb = DECISION_COLOR.get(str(v), RGBColor(0x26, 0x26, 0x26))
            else:
                run.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
    return table


# ============ بيانات مصدرها الحرفي DC.C.3.3.docx (Table 14) + DC.C.3.2.docx (Table 11) ============
ASSETS = [
    (1, 9, "بيانات طلبات ومعاملات المستفيدين", "الإبقاء كمقيّد"),
    (2, 10, "بيانات الخدمات الحكومية الرقمية", "إعادة التصنيف إلى عام"),
    (3, 11, "بيانات الموردين والمتعاقدين", "الإبقاء كمقيّد"),
    (4, 14, "بيانات التدريب والتطوير الوظيفي", "الإبقاء كمقيّد"),
    (5, 16, "بيانات الحوادث التقنية وطلبات الدعم", "الإبقاء كمقيّد"),
    (6, 21, "سجلات المراسلات والخطابات الرسمية", "الإبقاء كمقيّد"),
    (7, 22, "سجلات الاجتماعات واللجان الحوكمية", "الإبقاء كمقيّد"),
    (8, 23, "سجلات التقارير الإدارية الدورية", "إعادة التصنيف إلى عام"),
    (9, 25, "سجلات الشكاوى والمقترحات", "الإبقاء كمقيّد"),
    (10, 27, "سجلات التراخيص والتصاريح القانونية", "إعادة التصنيف إلى عام"),
    (11, 29, "وثائق السياسات والإجراءات الداخلية", "الإبقاء كمقيّد"),
    (12, 30, "سجلات التغييرات والتحديثات التقنية", "الإبقاء كمقيّد"),
]
TOTAL_LOW_IMPACT = len(ASSETS)
REMAINED_RESTRICTED = sum(1 for *_, d in ASSETS if d == "الإبقاء كمقيّد")
RECLASSIFIED = sum(1 for *_, d in ASSETS if d == "إعادة التصنيف إلى عام")
KPI_VALUE = round(REMAINED_RESTRICTED / TOTAL_LOW_IMPACT * 100)  # = 75

doc = docx.Document()
section = doc.sections[0]
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)
section.top_margin = Cm(1.5)
section.bottom_margin = Cm(1.5)

style = doc.styles["Normal"]
style.font.name = AR_FONT
style.font.size = Pt(10)
style.element.rPr.rFonts.set(qn("w:cs"), AR_FONT)

# ============ صفحة الغلاف / 1. بيانات الوثيقة ============
banner = doc.add_table(rows=1, cols=1)
banner.autofit = True
set_table_rtl(banner)
cell = banner.rows[0].cells[0]
set_cell_background(cell, "1F2A44")
cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
p = cell.paragraphs[0]
rtl(p)
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = p.add_run("سجل احتساب مؤشر مراجعة البيانات منخفضة الأثر (KPI-DC-03)")
run.font.size = Pt(18)
run.font.bold = True
run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
set_rtl_run(run)
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(6)

doc.add_paragraph()
add_paragraph_ar(doc, "منصة ذكية لإدارة وقياس الامتثال لمتطلبات المكتب الوطني لإدارة البيانات (NDMO)",
                  size=10.5, bold=True, color=GOLD)
add_heading_ar(doc, "سجل احتساب مؤشر مراجعة البيانات منخفضة الأثر (KPI-DC-03)", size=14)
add_paragraph_ar(doc, "Low-Impact Data Review KPI Calculation Record", size=11, italic=True, color=GOLD)

doc.add_paragraph()
add_heading_ar(doc, "1.  بيانات الوثيقة", size=11.5)
meta_table = doc.add_table(rows=0, cols=2)
meta_table.autofit = True
set_table_rtl(meta_table)
meta_table.columns[0].width = Cm(11.5)
meta_table.columns[1].width = Cm(5)
add_meta_row(meta_table, "الجهة", "هيئة الخدمات الحكومية الذكية (SGSA)")
add_meta_row(meta_table, "المنصة", "")
add_meta_row(meta_table, "الإطار المرجعي", "إطار حوكمة البيانات الوطني (NDMO)")
add_meta_row(meta_table, "رمز الدليل", "DC.C.4.1-E03")
add_meta_row(meta_table, "نوع الوثيقة", "Calculation Record")
add_meta_row(meta_table, "المتطلب المرتبط", "DC.C.4.1 — تقرير مراقبة عمليات تصنيف البيانات عبر مؤشرات الأداء الرئيسية")
add_meta_row(meta_table, "المؤشر المرتبط", "KPI-DC-03 — نسبة البيانات منخفضة الأثر المصنفة «مقيّد»")
add_meta_row(meta_table, "الإصدار", "1.0")
add_meta_row(meta_table, "التاريخ", "يونيو 2026")
add_meta_row(meta_table, "حالة الوثيقة", "جاهزة للاعتماد")

doc.add_page_break()

# ============ 2. نطاق القياس ============
add_heading_ar(doc, "2.  نطاق القياس")
meta2 = doc.add_table(rows=0, cols=2)
meta2.autofit = True
set_table_rtl(meta2)
meta2.columns[0].width = Cm(11.5)
meta2.columns[1].width = Cm(5)
add_meta_row(
    meta2, "نطاق الأصول محل المراجعة",
    "الأصول الاثنا عشر (12) المصنَّفة أصلاً «مقيّد» في تقرير تقييم الأثر (DC.C.3.2)، "
    "والتي خضعت جميعها دون استثناء لدراسة تفصيلية (تعارض نظامي + موازنة منافع/آثار) "
    "في تقرير DC.C.3.3."
)
add_meta_row(meta2, "فترة المراجعة", "الربع الثاني 2026 — كما ورد في تقرير DC.C.4.1")
add_meta_row(meta2, "تاريخ استخراج البيانات", "يونيو 2026 (تاريخ إصدار DC.C.3.2 وDC.C.3.3)")

doc.add_paragraph()

# ============ 3. مصدر البيانات ============
add_heading_ar(doc, "3.  مصدر البيانات")
add_bullet_ar(doc, "DC.C.3.2 — تقرير تقييم الأثر (تحديد الأصول الاثني عشر المصنَّفة أصلاً «مقيّد» ذات الأثر المنخفض).")
add_bullet_ar(doc, "DC.C.3.3 — تقرير تقييم البيانات منخفضة الأثر (القرار النهائي لكل أصل من الأصول الاثني عشر: الإبقاء كمقيّد أو إعادة التصنيف إلى عام).")

doc.add_paragraph()

# ============ 4. جدول البيانات المستخدمة ============
add_heading_ar(doc, "4.  جدول البيانات المستخدمة")
add_paragraph_ar(doc, "يعرض الجدول التالي القرار النهائي لكل أصل من الأصول الاثني عشر، كما ورد حرفياً في تقرير DC.C.3.3:")
rows = [(a, b, name, decision) for a, b, name, decision in ASSETS]
make_table(
    doc,
    ["م (DC.C.3.3)", "الرقم الأصلي (DC.C.3.1/3.2)", "اسم الأصل", "القرار النهائي (DC.C.3.3)"],
    rows,
    [Cm(2.2), Cm(2.8), Cm(6.5), Cm(4.5)],
    status_col=[3],
    font_size=8,
)

doc.add_paragraph()
add_paragraph_ar(
    doc,
    "ملاحظة تدقيق: لا يعرض هذا الجدول قائمة مستقلة «مُنشأة» للأصول التسعة الباقية على "
    "مقيّد؛ بل هو نسخ حرفي للجدول الكامل الوارد في DC.C.3.3 (القسم 7.2) لجميع الأصول "
    "الاثني عشر، والذي يتضمن أصلاً عمود القرار النهائي لكل صف. لم يُستخدم أي مصدر آخر "
    "غير DC.C.3.2 وDC.C.3.3 في هذا السجل.",
    size=8, italic=True, color=GREY,
)

doc.add_page_break()

# ============ 5. تطبيق المعادلة ============
add_heading_ar(doc, "5.  تطبيق المعادلة")
make_table(
    doc,
    ["البند", "القيمة"],
    [
        ("إجمالي الأصول منخفضة الأثر قبل المراجعة (DC.C.3.2)", str(TOTAL_LOW_IMPACT)),
        ("عدد الأصول التي بقيت مصنَّفة «مقيّد» بعد المراجعة (DC.C.3.3)", str(REMAINED_RESTRICTED)),
        ("عدد الأصول المُعاد تصنيفها إلى «عام» (DC.C.3.3)", str(RECLASSIFIED)),
        ("معادلة المؤشر", "عدد الأصول الباقية على مقيّد ÷ إجمالي الأصول منخفضة الأثر × 100"),
        ("تطبيق المعادلة", f"{REMAINED_RESTRICTED} ÷ {TOTAL_LOW_IMPACT} × 100"),
        ("القيمة النهائية للمؤشر", f"{KPI_VALUE}%"),
    ],
    [Cm(10.5), Cm(4.9)],
    status_col=[1],
    font_size=9,
)

doc.add_paragraph()

# ============ 6. النتيجة والتحقق مقابل DC.C.4.1 ============
add_heading_ar(doc, "6.  النتيجة والتحقق مقابل DC.C.4.1")
add_paragraph_ar(
    doc,
    f"القيمة الناتجة ({KPI_VALUE}%) مطابقة تماماً للقيمة المقيسة الواردة في بطاقة مؤشر "
    "KPI-DC-03 ضمن تقرير DC.C.4.1 (القسم 8 وبند 9.3: «75.0% (9 أصول بقيت مقيّدة من أصل "
    "12 أصلاً مقيّداً أولياً)»)، وللقيمة الواردة في جدول لوحة الأداء الإجمالية (القسم 10) "
    "وجدول التوافق مع متطلبات NDMO (القسم 14)."
)
add_bullet_ar(doc, "تم التحقق من أن الرقم التسلسلي الأصلي لكل أصل من الاثني عشر يطابق تصنيفه «مقيّد» الوارد في DC.C.3.2 (القسم 6.2)، دون أي أصل إضافي أو مفقود.")
add_bullet_ar(doc, "الأصول الثلاثة المُعاد تصنيفها (الأرقام الأصلية 10 و23 و27) تطابق تماماً الأصول الثلاثة الواردة في سجل DC.C.5.1 كأصول «أُعيد تصنيفها من مقيّد إلى عام».")

doc.add_paragraph()

# ============ 7. الاعتماد والتوقيع ============
add_heading_ar(doc, "7.  الاعتماد والتوقيع")
add_paragraph_ar(
    doc,
    "تمت مراجعة هذا السجل من مكتب إدارة البيانات، وتم التحقق من مطابقة البيانات الواردة "
    "فيه لمحتوى DC.C.3.2 وDC.C.3.3 دون أي إضافة أو حذف أو تعديل."
)
doc.add_paragraph()

approval_table = doc.add_table(rows=2, cols=4)
approval_table.style = "Table Grid"
set_table_rtl(approval_table)
headers = ["صاحب الصلاحية", "المسمى الوظيفي", "التوقيع", "التاريخ"]
hdr_cells = approval_table.rows[0].cells
for i, h in enumerate(headers):
    set_cell_background(hdr_cells[i], "1F2A44")
    hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = hdr_cells[i].paragraphs[0]
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(h)
    r.font.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    set_rtl_run(r)

vals = ["________________", "مدير مكتب إدارة البيانات", "________________", "________________"]
row_cells = approval_table.rows[1].cells
for i, v in enumerate(vals):
    p = row_cells[i].paragraphs[0]
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(v)
    r.font.size = Pt(9)
    set_rtl_run(r)

doc.save(OUT_PATH)
print("Saved:", OUT_PATH)
print("TOTAL_LOW_IMPACT:", TOTAL_LOW_IMPACT)
print("REMAINED_RESTRICTED:", REMAINED_RESTRICTED)
print("RECLASSIFIED:", RECLASSIFIED)
print("KPI_VALUE:", KPI_VALUE)
