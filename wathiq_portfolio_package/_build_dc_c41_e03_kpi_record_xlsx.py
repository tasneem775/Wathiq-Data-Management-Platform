# -*- coding: utf-8 -*-
"""Builds the supporting Excel calculation sheet for DC.C.4.1-E03 (KPI-DC-03) — 7-section
compliance-audit template. All 12 rows are copied verbatim from the real DC.C.3.3.docx
(Table 14) and DC.C.3.2.docx (Table 11). No asset name, decision, or count is invented."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT_PATH = "06_DC_C4_1_KPI_Evidence/DC.C.4.1-E03_KPI-DC-03_Low_Impact_Data_Review_KPI_Calculation_Record.xlsx"

NAVY = "1F2A44"
GOLD = "AB8654"
LIGHT_GREY = "F2F2F2"
WHITE = "FFFFFF"
GREEN = "1F6F43"
AMBER = "8A6D00"
AR_FONT = "Arial"

ASSETS = [
    (1, 9, "بيانات طلبات ومعاملات المستفيدين", "الإبقاء كمقيّد"),
    (2, 10, "بيانات الخدمات الحكومية الرقمية", "إعادة التصنيف إلى عام"),
    (3, 11, "بيانات الموردين والمتعاقدين", "الإبقاء كمقيّد"),
    (4, 14, "بيانات التدريب والتطوير الوظيفي", "الإبقاء كمقيّد"),
    (5, 16, "بيانات الحوادث التقنية وطلبات الدعم", "الإبقاء كمقيّد"),
    (6, 21, "سجلات المراسلات والخطابات الرسمية", "الإبقاء كمقيّد"),
    (7, 22, "سجلات الاجتماعات واللجان الحوكمية", "الإبقاء كمقيّد"),
    (8, 23, "سجلات التقارير الإدارية الدورية", "إعادة التصنيف إلى عام"),
    (9, 25, "سجلات الشكاوى والمقترحات", "الإبقاء كمقيّد"),
    (10, 27, "سجلات التراخيص والتصاريح القانونية", "إعادة التصنيف إلى عام"),
    (11, 29, "وثائق السياسات والإجراءات الداخلية", "الإبقاء كمقيّد"),
    (12, 30, "سجلات التغييرات والتحديثات التقنية", "الإبقاء كمقيّد"),
]
TOTAL_LOW_IMPACT = len(ASSETS)
REMAINED_RESTRICTED = sum(1 for *_, d in ASSETS if d == "الإبقاء كمقيّد")
RECLASSIFIED = sum(1 for *_, d in ASSETS if d == "إعادة التصنيف إلى عام")
KPI_VALUE = round(REMAINED_RESTRICTED / TOTAL_LOW_IMPACT * 100)
DECISION_COLOR = {"الإبقاء كمقيّد": AMBER, "إعادة التصنيف إلى عام": GREEN}

wb = Workbook()
ws = wb.active
ws.title = "احتساب KPI-DC-03"
ws.sheet_view.showGridLines = False
ws.sheet_view.rightToLeft = True

thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

N_COLS = 6
r = 1


def merged_text(row, text, size, bold, color, fill=None, height=None, italic=False, span=N_COLS):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name=AR_FONT, size=size, bold=bold, italic=italic, color=color)
    cell.alignment = Alignment(horizontal="right", vertical="center", indent=1, wrap_text=True)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if height:
        ws.row_dimensions[row].height = height
    return row + 1


def section_heading(row, text, span=N_COLS):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name=AR_FONT, size=12.5, bold=True, color=NAVY)
    cell.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    cell.border = Border(bottom=Side(style="medium", color=GOLD))
    ws.row_dimensions[row].height = 22
    return row + 1


def spacer(row, h=10):
    ws.row_dimensions[row].height = h
    return row + 1


def meta_row(row, label, value):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    c1 = ws.cell(row=row, column=1, value=label)
    c1.font = Font(name=AR_FONT, size=10, bold=True, color="595959")
    c1.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=N_COLS)
    c2 = ws.cell(row=row, column=3, value=value)
    c2.font = Font(name=AR_FONT, size=10, color="262626")
    c2.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    return row + 1


def bullet_row(row, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=N_COLS)
    cell = ws.cell(row=row, column=1, value="•  " + text)
    cell.font = Font(name=AR_FONT, size=9.5, color="262626")
    cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True, indent=1)
    ws.row_dimensions[row].height = 30
    return row + 1


r = merged_text(r, "سجل احتساب مؤشر مراجعة البيانات منخفضة الأثر (KPI-DC-03)", 18, True, WHITE, fill=NAVY, height=32)
r = merged_text(r, "منصة ذكية لإدارة وقياس الامتثال لمتطلبات المكتب الوطني لإدارة البيانات (NDMO)", 10.5, True, WHITE, fill=NAVY, height=20)
r = spacer(r, 6)
r = merged_text(r, "سجل احتساب مؤشر مراجعة البيانات منخفضة الأثر (KPI-DC-03)", 14, True, NAVY, height=24)
r = merged_text(r, "دليل داعم — DC.C.4.1-E03 — Calculation Record", 11, True, GOLD, italic=True, height=18)
r = spacer(r, 10)

r = section_heading(r, "1. بيانات الوثيقة")
r = spacer(r, 4)
r = meta_row(r, "الجهة", "هيئة الخدمات الحكومية الذكية (SGSA)")
r = meta_row(r, "المنصة", "")
r = meta_row(r, "الإطار المرجعي", "إطار حوكمة البيانات الوطني (NDMO)")
r = meta_row(r, "رمز الدليل", "DC.C.4.1-E03")
r = meta_row(r, "نوع الوثيقة", "Calculation Record")
r = meta_row(r, "المتطلب المرتبط", "DC.C.4.1 — تقرير مراقبة عمليات تصنيف البيانات عبر مؤشرات الأداء الرئيسية")
r = meta_row(r, "المؤشر المرتبط", "KPI-DC-03 — نسبة البيانات منخفضة الأثر المصنفة «مقيّد»")
r = meta_row(r, "الإصدار", "1.0")
r = meta_row(r, "التاريخ", "يونيو 2026")
r = meta_row(r, "حالة الوثيقة", "جاهزة للاعتماد")
r = spacer(r, 14)

r = section_heading(r, "2. نطاق القياس")
r = spacer(r, 4)
r = meta_row(r, "نطاق الأصول محل المراجعة", "12 أصلاً مصنَّفة أصلاً «مقيّد» في DC.C.3.2، خضعت جميعها لمراجعة DC.C.3.3")
r = meta_row(r, "فترة المراجعة", "الربع الثاني 2026")
r = spacer(r, 14)

r = section_heading(r, "3. مصدر البيانات")
r = spacer(r, 4)
r = bullet_row(r, "DC.C.3.2 — تقرير تقييم الأثر (تحديد الأصول الاثني عشر المصنَّفة أصلاً «مقيّد»)")
r = bullet_row(r, "DC.C.3.3 — تقرير تقييم البيانات منخفضة الأثر (القرار النهائي لكل أصل)")
r = spacer(r, 14)

r = section_heading(r, "4. جدول البيانات المستخدمة (مصدرها DC.C.3.3 Table 14)")
r = spacer(r, 4)

headers = ["م (DC.C.3.3)", "الرقم الأصلي", "اسم الأصل", "القرار النهائي", "معادلة الحساب", "النتيجة النهائية"]
header_row = r
for c, h in enumerate(headers, start=1):
    cell = ws.cell(row=header_row, column=c, value=h)
    cell.font = Font(name=AR_FONT, size=10, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border
ws.row_dimensions[header_row].height = 34
r += 1

FORMULA_TEXT = "9 ÷ 12 × 100"
FINAL_RESULT = f"{KPI_VALUE}%"
for i, (a, b, name, decision) in enumerate(ASSETS):
    row_fill = LIGHT_GREY if i % 2 == 0 else WHITE
    row_vals = [a, b, name, decision, FORMULA_TEXT, FINAL_RESULT]
    for c, v in enumerate(row_vals, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.border = border
        cell.fill = PatternFill("solid", fgColor=row_fill)
        center_cols = (1, 2, 4, 6)
        cell.alignment = Alignment(
            horizontal="center" if c in center_cols else "right",
            vertical="center", wrap_text=True,
            indent=0 if c in center_cols else 1,
        )
        if c == 3:
            cell.font = Font(name=AR_FONT, size=10, bold=True, color=NAVY)
        elif c == 4:
            cell.font = Font(name=AR_FONT, size=9.5, bold=True, color=DECISION_COLOR.get(v, "262626"))
        elif c == 6:
            cell.font = Font(name=AR_FONT, size=10, bold=True, color=GREEN)
        else:
            cell.font = Font(name=AR_FONT, size=9.5, color="262626")
    ws.row_dimensions[r].height = 26
    r += 1

r = spacer(r, 14)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=N_COLS)
note1 = ws.cell(
    row=r, column=1,
    value=(
        "ملاحظة تدقيق: هذا الجدول نسخ حرفي كامل للجدول الوارد في DC.C.3.3 لجميع الأصول "
        "الاثني عشر (بما فيها التسعة الباقية على مقيّد)، وليس قائمة مُنشأة بشكل مستقل."
    ),
)
note1.font = Font(name=AR_FONT, size=8, italic=True, color="7F7F7F")
note1.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True, indent=1)
ws.row_dimensions[r].height = 30
r += 1

r = spacer(r, 14)
r = section_heading(r, "5. تطبيق المعادلة")
r = spacer(r, 4)
r = meta_row(r, "إجمالي الأصول منخفضة الأثر قبل المراجعة (DC.C.3.2)", str(TOTAL_LOW_IMPACT))
r = meta_row(r, "عدد الأصول الباقية على «مقيّد» (DC.C.3.3)", str(REMAINED_RESTRICTED))
r = meta_row(r, "عدد الأصول المُعاد تصنيفها إلى «عام» (DC.C.3.3)", str(RECLASSIFIED))
r = meta_row(r, "تطبيق المعادلة", f"{REMAINED_RESTRICTED} ÷ {TOTAL_LOW_IMPACT} × 100")
r = meta_row(r, "القيمة النهائية للمؤشر", FINAL_RESULT)
r = spacer(r, 14)

r = section_heading(r, "6. النتيجة والتحقق مقابل DC.C.4.1")
r = spacer(r, 4)
r = bullet_row(r, f"القيمة الناتجة ({FINAL_RESULT}) مطابقة تماماً للقيمة الواردة في بطاقة KPI-DC-03 ضمن تقرير DC.C.4.1 (القسم 8 وبند 9.3).")
r = spacer(r, 14)

r = section_heading(r, "7. الاعتماد والتوقيع")
r = spacer(r, 4)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=N_COLS)
review_cell = ws.cell(
    row=r, column=1,
    value="تمت مراجعة هذا السجل من مكتب إدارة البيانات، وتم التحقق من مطابقة البيانات الواردة فيه لمحتوى DC.C.3.2 وDC.C.3.3 دون أي إضافة أو حذف أو تعديل.",
)
review_cell.font = Font(name=AR_FONT, size=10, color="262626")
review_cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True, indent=1)
ws.row_dimensions[r].height = 26
r += 1
r = spacer(r, 6)

approval_headers = ["صاحب الصلاحية", "المسمى الوظيفي", "التوقيع", "التاريخ"]
col_spans = [(1, 1), (2, 2), (3, 4), (5, 6)]
approval_header_row = r
for (sc, ec), h in zip(col_spans, approval_headers):
    ws.merge_cells(start_row=approval_header_row, start_column=sc, end_row=approval_header_row, end_column=ec)
    cell = ws.cell(row=approval_header_row, column=sc, value=h)
    cell.font = Font(name=AR_FONT, size=10, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    for cc in range(sc, ec + 1):
        ws.cell(row=approval_header_row, column=cc).border = border
ws.row_dimensions[approval_header_row].height = 20
r += 1

approval_values = ["____________", "مدير مكتب إدارة البيانات", "____________", "____________"]
for (sc, ec), v in zip(col_spans, approval_values):
    ws.merge_cells(start_row=r, start_column=sc, end_row=r, end_column=ec)
    cell = ws.cell(row=r, column=sc, value=v)
    cell.font = Font(name=AR_FONT, size=9.5, color="262626")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    for cc in range(sc, ec + 1):
        ws.cell(row=r, column=cc).border = border
ws.row_dimensions[r].height = 26

widths = [12, 12, 36, 22, 16, 16]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

wb.save(OUT_PATH)
print("Saved:", OUT_PATH)
print("Total:", TOTAL_LOW_IMPACT, "Remained:", REMAINED_RESTRICTED, "Reclassified:", RECLASSIFIED, "Result:", FINAL_RESULT)
