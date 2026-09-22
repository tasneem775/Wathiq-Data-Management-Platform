# -*- coding: utf-8 -*-
"""Builds a standalone compliance-audit Evidence Record for DC.C.4.1 (KPI-DC-04).

IMPORTANT — this is an Evidence Record, NOT a Calculation Record: the only data available
for KPI-DC-04 (نسبة الأصول التي خضعت للمراجعة والتعميد والتصديق) is the aggregate 4-stage
approval table already stated in DC.C.4.1.docx itself (Table 17 under section 9.4). No
per-asset review/endorsement/certification record with individual signatures exists in any
evidence_repository document (DC.C.5.1 provides only a single merged review-date column per
asset, not a 4-stage breakdown). This document therefore reproduces and documents the stated
aggregate figure without claiming an independent per-asset recomputation. DC.C.4.1 itself is
not modified.

Audit template (6 sections — "تطبيق المعادلة" is omitted per the brief, which restricts that
section to Calculation Records only):
1. بيانات الوثيقة  2. نطاق القياس  3. مصدر البيانات  4. جدول البيانات المستخدمة
5. النتيجة والتحقق مقابل DC.C.4.1  6. الاعتماد والتوقيع
"""
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_PATH = "06_DC_C4_1_KPI_Evidence/DC.C.4.1-E04_KPI-DC-04_Review_Endorsement_Approval_KPI_Evidence_Record.docx"

NAVY = RGBColor(0x1F, 0x2A, 0x44)
GOLD = RGBColor(0xAB, 0x8D, 0x54)
GREY = RGBColor(0x59, 0x59, 0x59)
GREEN = RGBColor(0x1F, 0x6F, 0x43)
AR_FONT = "Arial"


def rtl(p):
    b = OxmlElement("w:bidi")
    b.set(qn("w:val"), "1")
    p._p.get_or_add_pPr().append(b)


def set_rtl_run(run):
    pr = run._element.get_or_add_rPr()
    pr.get_or_add_rFonts().set(qn("w:cs"), AR_FONT)
    e = OxmlElement("w:rtl")
    e.set(qn("w:val"), "1")
    pr.append(e)


def set_table_rtl(table):
    tbl_pr = table._tbl.tblPr
    bv = OxmlElement("w:bidiVisual")
    bv.set(qn("w:val"), "1")
    tbl_pr.append(bv)


def set_cell_background(cell, color_hex):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color_hex)
    cell._tc.get_or_add_tcPr().append(shd)


def add_meta_row(table, label, value):
    row = table.add_row()
    p1 = row.cells[1].paragraphs[0]
    rtl(p1)
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p1.add_run(label)
    r1.font.bold = True
    r1.font.size = Pt(9)
    r1.font.color.rgb = GREY
    set_rtl_run(r1)

    p2 = row.cells[0].paragraphs[0]
    rtl(p2)
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = p2.add_run(value)
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
    set_rtl_run(r2)


def add_paragraph_ar(doc, text, size=9.5, bold=False, italic=False, color=RGBColor(0x26, 0x26, 0x26)):
    p = doc.add_paragraph()
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    set_rtl_run(r)
    return p


def add_heading_ar(doc, text, size=13):
    p = doc.add_paragraph()
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = NAVY
    set_rtl_run(r)
    return p


def add_bullet_ar(doc, text, size=9.5):
    p = doc.add_paragraph()
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("•  " + text)
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
    set_rtl_run(r)
    return p


def make_table(doc, headers, rows, col_widths, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_rtl(table)
    for i, w in enumerate(col_widths):
        table.columns[i].width = w

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        set_cell_background(hdr_cells[i], "1F2A44")
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = hdr_cells[i].paragraphs[0]
        rtl(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_rtl_run(r)

    for idx, row_vals in enumerate(rows):
        row = table.add_row()
        for c, v in enumerate(row_vals):
            cell = row.cells[c]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if idx % 2 == 0:
                set_cell_background(cell, "F2F2F2")
            p = cell.paragraphs[0]
            rtl(p)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(v))
            run.font.size = Pt(font_size)
            set_rtl_run(run)
            if c == 0:
                run.font.bold = True
                run.font.color.rgb = NAVY
            elif c == 1:
                run.font.bold = True
                run.font.color.rgb = GREEN
            else:
                run.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
    return table


# ============ بيانات مصدرها الحرفي DC.C.4.1.docx (Table 17، بند 9.4) ============
STAGES = [
    ("المراجعة الأولية من فريق التصنيف", 30, "100%", "مارس 2026", "فريق مكتب إدارة البيانات"),
    ("اعتماد لجنة حوكمة البيانات", 30, "100%", "أبريل 2026", "لجنة حوكمة البيانات"),
    ("التصديق من مدير مكتب إدارة البيانات", 30, "100%", "مايو 2026", "مدير مكتب إدارة البيانات"),
    ("الاعتماد النهائي من مدير عام الهيئة", 30, "100%", "يونيو 2026", "مدير عام الهيئة"),
]
TOTAL_ASSETS = 30
KPI_VALUE = "100%"

doc = docx.Document()
section = doc.sections[0]
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)
section.top_margin = Cm(1.5)
section.bottom_margin = Cm(1.5)

style = doc.styles["Normal"]
style.font.name = AR_FONT
style.font.size = Pt(10)
style.element.rPr.rFonts.set(qn("w:cs"), AR_FONT)

# ============ صفحة الغلاف / 1. بيانات الوثيقة ============
banner = doc.add_table(rows=1, cols=1)
banner.autofit = True
set_table_rtl(banner)
cell = banner.rows[0].cells[0]
set_cell_background(cell, "1F2A44")
cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
p = cell.paragraphs[0]
rtl(p)
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = p.add_run("سجل إثبات مؤشر المراجعة والتعميد والتصديق (KPI-DC-04)")
run.font.size = Pt(18)
run.font.bold = True
run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
set_rtl_run(run)
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(6)

doc.add_paragraph()
add_paragraph_ar(doc, "منصة ذكية لإدارة وقياس الامتثال لمتطلبات المكتب الوطني لإدارة البيانات (NDMO)",
                  size=10.5, bold=True, color=GOLD)
add_heading_ar(doc, "سجل إثبات مؤشر المراجعة والتعميد والتصديق (KPI-DC-04)", size=14)
add_paragraph_ar(doc, "Review, Endorsement & Approval KPI Evidence Record", size=11, italic=True, color=GOLD)

doc.add_paragraph()
add_heading_ar(doc, "1.  بيانات الوثيقة", size=11.5)
meta_table = doc.add_table(rows=0, cols=2)
meta_table.autofit = True
set_table_rtl(meta_table)
meta_table.columns[0].width = Cm(11.5)
meta_table.columns[1].width = Cm(5)
add_meta_row(meta_table, "الجهة", "هيئة الخدمات الحكومية الذكية (SGSA)")
add_meta_row(meta_table, "المنصة", "")
add_meta_row(meta_table, "الإطار المرجعي", "إطار حوكمة البيانات الوطني (NDMO)")
add_meta_row(meta_table, "رمز الدليل", "DC.C.4.1-E04")
add_meta_row(meta_table, "نوع الوثيقة", "Evidence Record")
add_meta_row(meta_table, "المتطلب المرتبط", "DC.C.4.1 — تقرير مراقبة عمليات تصنيف البيانات عبر مؤشرات الأداء الرئيسية")
add_meta_row(meta_table, "المؤشر المرتبط", "KPI-DC-04 — نسبة الأصول التي خضعت للمراجعة والتعميد والتصديق")
add_meta_row(meta_table, "الإصدار", "1.0")
add_meta_row(meta_table, "التاريخ", "يونيو 2026")
add_meta_row(meta_table, "حالة الوثيقة", "جاهزة للاعتماد")

doc.add_page_break()

# ============ 2. نطاق القياس ============
add_heading_ar(doc, "2.  نطاق القياس")
meta2 = doc.add_table(rows=0, cols=2)
meta2.autofit = True
set_table_rtl(meta2)
meta2.columns[0].width = Cm(11.5)
meta2.columns[1].width = Cm(5)
add_meta_row(meta2, "عدد الأصول", "30 أصلاً (إجمالي الأصول البيانية المصنَّفة وفق DC.C.3.1)")
add_meta_row(meta2, "فترة الاعتماد", "مارس 2026 – يونيو 2026 — أربع مراحل متتابعة كما ورد في DC.C.4.1")
add_meta_row(meta2, "تاريخ استخراج البيانات", "يونيو 2026 (تاريخ إصدار DC.C.4.1)")

doc.add_paragraph()

# ============ 3. مصدر البيانات ============
add_heading_ar(doc, "3.  مصدر البيانات")
add_bullet_ar(doc, "تقرير DC.C.4.1 نفسه (القسم 9.4 وجدول مراحل الاعتماد) — المصدر الوحيد المتاح لبيانات هذا المؤشر بمستوى المراحل الأربع.")
add_bullet_ar(doc, "DC.C.5.1 — سجل البيانات، يثبت لكل أصل من الأصول الثلاثين تاريخ مراجعة تصنيف واحداً (يونيو 2026)، دون تفصيل مراحل التعميد والتصديق الأربع بشكل منفصل لكل أصل.")

doc.add_paragraph()

# ============ 4. جدول البيانات المستخدمة ============
add_heading_ar(doc, "4.  جدول البيانات المستخدمة")
add_paragraph_ar(doc, "استُخدمت فقط البيانات الإجمالية التالية، كما وردت حرفياً في تقرير DC.C.4.1 (القسم 9.4):")
rows = [(stage, str(n), pct, date, entity) for stage, n, pct, date, entity in STAGES]
make_table(
    doc,
    ["مرحلة الاعتماد", "عدد الأصول المكتملة", "النسبة", "تاريخ الاكتمال", "الجهة المعتمِدة"],
    rows,
    [Cm(5.0), Cm(2.6), Cm(1.8), Cm(2.6), Cm(4.5)],
    font_size=8.5,
)

doc.add_paragraph()
add_paragraph_ar(
    doc,
    "ملاحظة تدقيق: لا يتضمن هذا السجل سجلاً فردياً لكل أصل من الأصول الثلاثين، ولا "
    "توقيعات فردية، ولا أسماء معتمدين لكل أصل على حدة، لأن هذا المستوى من التفصيل غير "
    "متوفر في أي مصدر رسمي ضمن مستودع الأدلة. لا ينطبق قسم «تطبيق المعادلة» على هذا "
    "السجل: لا يتوفر سجل فردي لكل أصل يتيح إعادة احتساب النسبة (100%) من بيانات خام؛ "
    "يقتصر هذا السجل على إثبات القيمة الإجمالية المذكورة في DC.C.4.1 دون إعادة اشتقاقها.",
    size=8, italic=True, color=GREY,
)

doc.add_page_break()

# ============ 5. النتيجة والتحقق مقابل DC.C.4.1 ============
add_heading_ar(doc, "5.  النتيجة والتحقق مقابل DC.C.4.1")
add_paragraph_ar(
    doc,
    f"القيمة الموثَّقة ({KPI_VALUE}) مطابقة تماماً للقيمة الواردة في بطاقة مؤشر KPI-DC-04 "
    "ضمن تقرير DC.C.4.1 (القسم 8 وبند 9.4)، وللقيمة الواردة في جدول لوحة الأداء الإجمالية "
    "(القسم 10) وجدول التوافق مع متطلبات NDMO (القسم 14)."
)
add_bullet_ar(doc, "جدول مراحل الاعتماد الأربع (المراجعة الأولية، اعتماد اللجنة، التصديق، الاعتماد النهائي) هو نفسه الجدول الوارد حرفياً في تقرير DC.C.4.1، ولم تُضَف أو تُحذف أي مرحلة.")
add_bullet_ar(doc, "لم تُخترع أي توقيعات فردية أو أسماء معتمدين لكل أصل؛ الجهات المعتمِدة الأربع المذكورة في الجدول هي نفسها المذكورة في مصدرها فقط بصفتها الوظيفية (لا بأسماء أشخاص).")
add_bullet_ar(doc, "هذا السجل Evidence Record وليس Calculation Record، لأن البيانات المتاحة إجمالية فقط (مستوى المرحلة) ولا تتوفر بيانات خام فردية لكل أصل تتيح إعادة الاحتساب من الصفر.")

doc.add_paragraph()

# ============ 6. الاعتماد والتوقيع ============
add_heading_ar(doc, "6.  الاعتماد والتوقيع")
add_paragraph_ar(
    doc,
    "تمت مراجعة هذا السجل من مكتب إدارة البيانات، وتم التحقق من مطابقة البيانات الواردة "
    "فيه لمحتوى تقرير DC.C.4.1 دون أي إضافة أو حذف أو تعديل."
)
doc.add_paragraph()

approval_table = doc.add_table(rows=2, cols=4)
approval_table.style = "Table Grid"
set_table_rtl(approval_table)
headers = ["صاحب الصلاحية", "المسمى الوظيفي", "التوقيع", "التاريخ"]
hdr_cells = approval_table.rows[0].cells
for i, h in enumerate(headers):
    set_cell_background(hdr_cells[i], "1F2A44")
    hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = hdr_cells[i].paragraphs[0]
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(h)
    r.font.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    set_rtl_run(r)

vals = ["________________", "مدير مكتب إدارة البيانات", "________________", "________________"]
row_cells = approval_table.rows[1].cells
for i, v in enumerate(vals):
    p = row_cells[i].paragraphs[0]
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(v)
    r.font.size = Pt(9)
    set_rtl_run(r)

doc.save(OUT_PATH)
print("Saved:", OUT_PATH)
print("Stages:", len(STAGES))
print("KPI_VALUE:", KPI_VALUE)
