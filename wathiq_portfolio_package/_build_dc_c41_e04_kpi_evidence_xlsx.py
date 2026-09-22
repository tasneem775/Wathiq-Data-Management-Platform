# -*- coding: utf-8 -*-
"""Builds the supporting Excel evidence sheet for DC.C.4.1-E04 (KPI-DC-04) — 6-section
compliance-audit template (Evidence Record; "تطبيق المعادلة" omitted per brief). The 4-stage
aggregate table is copied verbatim from DC.C.4.1.docx itself (section 9.4). No per-asset
data, no invented signatures."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT_PATH = "06_DC_C4_1_KPI_Evidence/DC.C.4.1-E04_KPI-DC-04_Review_Endorsement_Approval_KPI_Evidence_Record.xlsx"

NAVY = "1F2A44"
GOLD = "AB8654"
LIGHT_GREY = "F2F2F2"
WHITE = "FFFFFF"
GREEN = "1F6F43"
AR_FONT = "Arial"

STAGES = [
    ("المراجعة الأولية من فريق التصنيف", 30, "100%", "مارس 2026", "فريق مكتب إدارة البيانات"),
    ("اعتماد لجنة حوكمة البيانات", 30, "100%", "أبريل 2026", "لجنة حوكمة البيانات"),
    ("التصديق من مدير مكتب إدارة البيانات", 30, "100%", "مايو 2026", "مدير مكتب إدارة البيانات"),
    ("الاعتماد النهائي من مدير عام الهيئة", 30, "100%", "يونيو 2026", "مدير عام الهيئة"),
]
KPI_VALUE = "100%"

wb = Workbook()
ws = wb.active
ws.title = "إثبات KPI-DC-04"
ws.sheet_view.showGridLines = False
ws.sheet_view.rightToLeft = True

thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

N_COLS = 6
r = 1


def merged_text(row, text, size, bold, color, fill=None, height=None, italic=False, span=N_COLS):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name=AR_FONT, size=size, bold=bold, italic=italic, color=color)
    cell.alignment = Alignment(horizontal="right", vertical="center", indent=1, wrap_text=True)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if height:
        ws.row_dimensions[row].height = height
    return row + 1


def section_heading(row, text, span=N_COLS):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name=AR_FONT, size=12.5, bold=True, color=NAVY)
    cell.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    cell.border = Border(bottom=Side(style="medium", color=GOLD))
    ws.row_dimensions[row].height = 22
    return row + 1


def spacer(row, h=10):
    ws.row_dimensions[row].height = h
    return row + 1


def meta_row(row, label, value):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    c1 = ws.cell(row=row, column=1, value=label)
    c1.font = Font(name=AR_FONT, size=10, bold=True, color="595959")
    c1.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=N_COLS)
    c2 = ws.cell(row=row, column=3, value=value)
    c2.font = Font(name=AR_FONT, size=10, color="262626")
    c2.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    return row + 1


def bullet_row(row, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=N_COLS)
    cell = ws.cell(row=row, column=1, value="•  " + text)
    cell.font = Font(name=AR_FONT, size=9.5, color="262626")
    cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True, indent=1)
    ws.row_dimensions[row].height = 30
    return row + 1


r = merged_text(r, "سجل إثبات مؤشر المراجعة والتعميد والتصديق (KPI-DC-04)", 18, True, WHITE, fill=NAVY, height=32)
r = merged_text(r, "منصة ذكية لإدارة وقياس الامتثال لمتطلبات المكتب الوطني لإدارة البيانات (NDMO)", 10.5, True, WHITE, fill=NAVY, height=20)
r = spacer(r, 6)
r = merged_text(r, "سجل إثبات مؤشر المراجعة والتعميد والتصديق (KPI-DC-04)", 14, True, NAVY, height=24)
r = merged_text(r, "دليل داعم — DC.C.4.1-E04 — Evidence Record", 11, True, GOLD, italic=True, height=18)
r = spacer(r, 10)

r = section_heading(r, "1. بيانات الوثيقة")
r = spacer(r, 4)
r = meta_row(r, "الجهة", "هيئة الخدمات الحكومية الذكية (SGSA)")
r = meta_row(r, "المنصة", "")
r = meta_row(r, "الإطار المرجعي", "إطار حوكمة البيانات الوطني (NDMO)")
r = meta_row(r, "رمز الدليل", "DC.C.4.1-E04")
r = meta_row(r, "نوع الوثيقة", "Evidence Record")
r = meta_row(r, "المتطلب المرتبط", "DC.C.4.1 — تقرير مراقبة عمليات تصنيف البيانات عبر مؤشرات الأداء الرئيسية")
r = meta_row(r, "المؤشر المرتبط", "KPI-DC-04 — نسبة الأصول التي خضعت للمراجعة والتعميد والتصديق")
r = meta_row(r, "الإصدار", "1.0")
r = meta_row(r, "التاريخ", "يونيو 2026")
r = meta_row(r, "حالة الوثيقة", "جاهزة للاعتماد")
r = spacer(r, 14)

r = section_heading(r, "2. نطاق القياس")
r = spacer(r, 4)
r = meta_row(r, "عدد الأصول", "30 أصلاً")
r = meta_row(r, "فترة الاعتماد", "مارس 2026 – يونيو 2026 — أربع مراحل متتابعة")
r = spacer(r, 14)

r = section_heading(r, "3. مصدر البيانات")
r = spacer(r, 4)
r = bullet_row(r, "تقرير DC.C.4.1 نفسه (القسم 9.4 وجدول مراحل الاعتماد)")
r = bullet_row(r, "DC.C.5.1 — سجل البيانات (تاريخ مراجعة تصنيف واحد لكل أصل، دون تفصيل 4 مراحل)")
r = spacer(r, 14)

r = section_heading(r, "4. جدول البيانات المستخدمة (مصدرها DC.C.4.1، القسم 9.4)")
r = spacer(r, 4)

headers = ["مرحلة الاعتماد", "عدد الأصول", "النسبة", "تاريخ الاكتمال", "الجهة المعتمِدة"]
header_row = r
col_spans_h = [(1, 1), (2, 2), (3, 3), (4, 4), (5, 6)]
for (sc, ec), h in zip(col_spans_h, headers):
    ws.merge_cells(start_row=header_row, start_column=sc, end_row=header_row, end_column=ec)
    cell = ws.cell(row=header_row, column=sc, value=h)
    cell.font = Font(name=AR_FONT, size=10, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cc in range(sc, ec + 1):
        ws.cell(row=header_row, column=cc).border = border
ws.row_dimensions[header_row].height = 30
r += 1

for i, (stage, n, pct, date, entity) in enumerate(STAGES):
    row_fill = LIGHT_GREY if i % 2 == 0 else WHITE
    vals = [stage, str(n), pct, date, entity]
    for (sc, ec), v in zip(col_spans_h, vals):
        ws.merge_cells(start_row=r, start_column=sc, end_row=r, end_column=ec)
        cell = ws.cell(row=r, column=sc, value=v)
        cell.border = border
        cell.fill = PatternFill("solid", fgColor=row_fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for cc in range(sc, ec + 1):
            ws.cell(row=r, column=cc).fill = PatternFill("solid", fgColor=row_fill)
            ws.cell(row=r, column=cc).border = border
        if sc == 1:
            cell.font = Font(name=AR_FONT, size=10, bold=True, color=NAVY)
        elif sc == 2:
            cell.font = Font(name=AR_FONT, size=10, bold=True, color=GREEN)
        else:
            cell.font = Font(name=AR_FONT, size=9.5, color="262626")
    ws.row_dimensions[r].height = 28
    r += 1

r = spacer(r, 14)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=N_COLS)
note1 = ws.cell(
    row=r, column=1,
    value=(
        "ملاحظة تدقيق: لا يتضمن هذا الجدول سجلاً فردياً لكل أصل من الأصول الثلاثين، ولا "
        "توقيعات فردية، ولا أسماء معتمدين لكل أصل على حدة. لا ينطبق قسم «تطبيق المعادلة» "
        "على هذا السجل لعدم توفر بيانات خام فردية لكل أصل."
    ),
)
note1.font = Font(name=AR_FONT, size=8, italic=True, color="7F7F7F")
note1.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True, indent=1)
ws.row_dimensions[r].height = 34
r += 1

r = spacer(r, 14)
r = section_heading(r, "5. النتيجة والتحقق مقابل DC.C.4.1")
r = spacer(r, 4)
r = bullet_row(r, f"القيمة الموثَّقة ({KPI_VALUE}) مطابقة تماماً للقيمة الواردة في بطاقة KPI-DC-04 ضمن تقرير DC.C.4.1 (القسم 8 وبند 9.4).")
r = bullet_row(r, "هذا السجل Evidence Record وليس Calculation Record لعدم توفر بيانات خام فردية لكل أصل.")
r = spacer(r, 14)

r = section_heading(r, "6. الاعتماد والتوقيع")
r = spacer(r, 4)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=N_COLS)
review_cell = ws.cell(
    row=r, column=1,
    value="تمت مراجعة هذا السجل من مكتب إدارة البيانات، وتم التحقق من مطابقة البيانات الواردة فيه لمحتوى تقرير DC.C.4.1 دون أي إضافة أو حذف أو تعديل.",
)
review_cell.font = Font(name=AR_FONT, size=10, color="262626")
review_cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True, indent=1)
ws.row_dimensions[r].height = 26
r += 1
r = spacer(r, 6)

approval_headers = ["صاحب الصلاحية", "المسمى الوظيفي", "التوقيع", "التاريخ"]
col_spans = [(1, 1), (2, 2), (3, 4), (5, 6)]
approval_header_row = r
for (sc, ec), h in zip(col_spans, approval_headers):
    ws.merge_cells(start_row=approval_header_row, start_column=sc, end_row=approval_header_row, end_column=ec)
    cell = ws.cell(row=approval_header_row, column=sc, value=h)
    cell.font = Font(name=AR_FONT, size=10, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    for cc in range(sc, ec + 1):
        ws.cell(row=approval_header_row, column=cc).border = border
ws.row_dimensions[approval_header_row].height = 20
r += 1

approval_values = ["____________", "مدير مكتب إدارة البيانات", "____________", "____________"]
for (sc, ec), v in zip(col_spans, approval_values):
    ws.merge_cells(start_row=r, start_column=sc, end_row=r, end_column=ec)
    cell = ws.cell(row=r, column=sc, value=v)
    cell.font = Font(name=AR_FONT, size=9.5, color="262626")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    for cc in range(sc, ec + 1):
        ws.cell(row=r, column=cc).border = border
ws.row_dimensions[r].height = 26

widths = [26, 12, 10, 16, 26, 10]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

wb.save(OUT_PATH)
print("Saved:", OUT_PATH)
print("Stages:", len(STAGES), "KPI_VALUE:", KPI_VALUE)
