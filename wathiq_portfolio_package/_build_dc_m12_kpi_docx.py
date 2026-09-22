# -*- coding: utf-8 -*-
"""Builds the 7 DOCX KPI records for DC.M.12 (06_DC_M12_KPI_Evidence/).

Standalone supporting records, independent from any existing evidence_repository file.
DC.M.12.docx is used ONLY to verify each record's final result (never as a raw-data source).
Raw data comes exclusively from DC.C.3.4, DC.C.3.5, DC.C.5.1 — see _dc_m12_kpi_data.py.
"""
import os
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from _dc_m12_kpi_data import RECORDS, DOC_META_COMMON

OUT_DIR = "06_DC_M12_KPI_Evidence"

NAVY = RGBColor(0x1F, 0x2A, 0x44)
GOLD = RGBColor(0xAB, 0x8D, 0x54)
GREY = RGBColor(0x59, 0x59, 0x59)
GREEN = RGBColor(0x1F, 0x6F, 0x43)
AMBER = RGBColor(0x8A, 0x6D, 0x00)
AR_FONT = "Arial"


# ترتيب العناصر داخل w:pPr/w:tblPr/w:sectPr مُلزم بمخطط OOXML (CT_PPrBase/CT_TblPrBase/
# CT_SectPrBase). إلحاق عنصر جديد بـ append() الخام في آخر القائمة قد يضعه بعد عنصر يجب أن
# يسبقه فعلياً (مثل tblW/tblLook قبل bidiVisual) — وهذا لا يكسر صحة XML لكنه يجعل Word
# الحقيقي يتجاهل العنصر بصمت عند العرض (يبقى المستند يظهر LTR رغم وجود bidi/bidiVisual في
# الملف). لذلك تُستخدم هنا insert_element_before بدل append() لضمان أن Word يطبّق RTL فعلياً.
_PPR_BIDI_SUCCESSORS = (
    "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
    "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)
_TBLPR_BIDIVISUAL_SUCCESSORS = (
    "w:tblStyleRowBandSize", "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing",
    "w:tblInd", "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
    "w:tblCaption", "w:tblDescription", "w:tblPrChange",
)
_SECTPR_BIDI_SUCCESSORS = ("w:rtlGutter", "w:docGrid", "w:printerSettings", "w:sectPrChange")


def rtl(p):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:bidi")
    pPr.insert_element_before(b, *_PPR_BIDI_SUCCESSORS)


def set_rtl_run(run):
    pr = run._element.get_or_add_rPr()
    pr.get_or_add_rFonts().set(qn("w:cs"), AR_FONT)
    e = OxmlElement("w:rtl")
    e.set(qn("w:val"), "1")
    pr.append(e)


def set_table_rtl(table):
    tbl_pr = table._tbl.tblPr
    bv = OxmlElement("w:bidiVisual")
    tbl_pr.insert_element_before(bv, *_TBLPR_BIDIVISUAL_SUCCESSORS)


def set_cell_background(cell, color_hex):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color_hex)
    cell._tc.get_or_add_tcPr().append(shd)


def set_section_rtl(section):
    """Marks the whole page/section as right-to-left (column order, reading order),
    so the document never renders as an LTR page — formatting only, no content change."""
    sectPr = section._sectPr
    bidi = OxmlElement("w:bidi")
    sectPr.insert_element_before(bidi, *_SECTPR_BIDI_SUCCESSORS)


def set_style_rtl(style, size_pt=11.5):
    """Sets the Normal style's default paragraph direction to RTL and base body size,
    so any paragraph without explicit overrides still renders RTL/right-aligned."""
    pPr = style.element.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    pPr.insert_element_before(bidi, *_PPR_BIDI_SUCCESSORS)
    style.font.size = Pt(size_pt)


def add_page_number_field(paragraph):
    run = paragraph.add_run()
    set_rtl_run(run)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    return run


def add_header_footer(doc, doc_code):
    """Running header (برند وثيق + الشعار الفرعي) وتذييل (رمز الدليل + رقم الصفحة) على كل
    صفحة — هوية بصرية موحدة، بلا أي شعارات أو أسماء جهات جديدة."""
    section = doc.sections[0]
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False

    hdr_p1 = section.header.paragraphs[0]
    rtl(hdr_p1)
    hdr_p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = hdr_p1.add_run("هيئة الخدمات الحكومية الذكية (SGSA)")
    r1.font.bold = True
    r1.font.size = Pt(12)
    r1.font.color.rgb = NAVY
    set_rtl_run(r1)

    hdr_p2 = section.header.add_paragraph()
    rtl(hdr_p2)
    hdr_p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = hdr_p2.add_run("منصة ذكية لإدارة وقياس الامتثال لمتطلبات المكتب الوطني لإدارة البيانات (NDMO)")
    r2.font.size = Pt(9)
    r2.font.color.rgb = GOLD
    set_rtl_run(r2)
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "AB8654")
    pBdr.append(bottom)
    hdr_p2.paragraph_format.element.get_or_add_pPr().append(pBdr)

    ftr_p = section.footer.paragraphs[0]
    rtl(ftr_p)
    ftr_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = ftr_p.add_run(f"{doc_code}    |    صفحة ")
    r3.font.size = Pt(8.5)
    r3.font.color.rgb = GREY
    set_rtl_run(r3)
    add_page_number_field(ftr_p)


def add_meta_row(table, label, value):
    row = table.add_row()
    p1 = row.cells[1].paragraphs[0]
    rtl(p1)
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p1.add_run(label)
    r1.font.bold = True
    r1.font.size = Pt(10.5)
    r1.font.color.rgb = GREY
    set_rtl_run(r1)

    p2 = row.cells[0].paragraphs[0]
    rtl(p2)
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = p2.add_run(value)
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
    set_rtl_run(r2)


def add_paragraph_ar(doc, text, size=11.5, bold=False, italic=False, color=RGBColor(0x26, 0x26, 0x26)):
    p = doc.add_paragraph()
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    set_rtl_run(r)
    return p


def add_heading_ar(doc, text, size=15):
    p = doc.add_paragraph()
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = NAVY
    set_rtl_run(r)
    return p


def add_bullet_ar(doc, text, size=11):
    p = doc.add_paragraph()
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("•  " + text)
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
    set_rtl_run(r)
    return p


def make_table(doc, headers, rows, col_widths, status_col=None, font_size=10):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_rtl(table)
    for i, w in enumerate(col_widths):
        table.columns[i].width = w

    # يمنع انقسام صف العناوين بين الصفحات ويكرره كرأس جدول عند الامتداد لصفحة تالية
    hdr_trPr = table.rows[0]._tr.get_or_add_trPr()
    hdr_trPr.append(OxmlElement("w:tblHeader"))

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        set_cell_background(hdr_cells[i], "1F2A44")
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = hdr_cells[i].paragraphs[0]
        rtl(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_rtl_run(r)

    status_col = status_col or []
    for idx, row_vals in enumerate(rows):
        row = table.add_row()
        for c, v in enumerate(row_vals):
            cell = row.cells[c]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if idx % 2 == 0:
                set_cell_background(cell, "F2F2F2")
            p = cell.paragraphs[0]
            rtl(p)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c in status_col else WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(str(v))
            run.font.size = Pt(font_size)
            set_rtl_run(run)
            if c == 0:
                run.font.bold = True
                run.font.color.rgb = NAVY
            elif c in status_col:
                run.font.bold = True
                run.font.color.rgb = GREEN if str(v) in ("تثبيت التصنيف", "مُنشور", "محقق") else AMBER
            else:
                run.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
    return table


def build_record(rec):
    out_path = os.path.join(OUT_DIR, rec["docx_name"] + ".docx")
    is_calc = rec["record_type"] == "calculation"

    doc = docx.Document()
    section = doc.sections[0]
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)
    set_section_rtl(section)

    style = doc.styles["Normal"]
    style.font.name = AR_FONT
    style.element.rPr.rFonts.set(qn("w:cs"), AR_FONT)
    set_style_rtl(style, size_pt=11.5)

    add_header_footer(doc, rec["doc_code"])

    # ---- غلاف ----
    banner = doc.add_table(rows=1, cols=1)
    banner.autofit = True
    set_table_rtl(banner)
    cell = banner.rows[0].cells[0]
    set_cell_background(cell, "1F2A44")
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("هيئة الخدمات الحكومية الذكية (SGSA)")
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    set_rtl_run(run)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)

    doc.add_paragraph()
    add_paragraph_ar(doc, "منصة ذكية لإدارة وقياس الامتثال لمتطلبات المكتب الوطني لإدارة البيانات (NDMO)",
                      size=11.5, bold=True, color=GOLD)
    record_type_ar = "سجل احتساب" if is_calc else "سجل إثبات"
    add_heading_ar(doc, f"{record_type_ar} مؤشر «{rec['kpi_name_ar']}» ({rec['kpi_code']})", size=17)
    suffix_en = "KPI Calculation Record" if is_calc else "KPI Evidence Record"
    add_paragraph_ar(doc, f"{rec['kpi_name_en']} — {suffix_en}", size=12.5, italic=True, color=GOLD)

    doc.add_paragraph()

    # ---- 1) بيانات الوثيقة ----
    meta_table = doc.add_table(rows=0, cols=2)
    meta_table.autofit = True
    set_table_rtl(meta_table)
    meta_table.columns[0].width = Cm(11)
    meta_table.columns[1].width = Cm(5)
    add_meta_row(meta_table, "الجهة", DOC_META_COMMON["entity"])
    add_meta_row(meta_table, "المنصة", DOC_META_COMMON["platform"])
    add_meta_row(meta_table, "الإطار المرجعي", DOC_META_COMMON["framework"])
    add_meta_row(meta_table, "رمز الدليل", rec["doc_code"])
    add_meta_row(meta_table, "المتطلب المرتبط", DOC_META_COMMON["related_requirement"])
    add_meta_row(meta_table, "المؤشر المرتبط", f"{rec['kpi_code']} — {rec['kpi_name_ar']}")
    add_meta_row(meta_table, "نوع السجل", "سجل احتساب (Calculation Record)" if is_calc else "سجل إثبات (Evidence Record)")
    add_meta_row(meta_table, "الإصدار", DOC_META_COMMON["version"])
    add_meta_row(meta_table, "التاريخ", DOC_META_COMMON["date"])
    add_meta_row(meta_table, "حالة الوثيقة", DOC_META_COMMON["status"])

    doc.add_page_break()

    # ---- 2) نطاق القياس ----
    add_heading_ar(doc, "1.  نطاق القياس")
    add_paragraph_ar(doc, rec["scope"])
    doc.add_paragraph()

    # ---- 3) مصادر البيانات ----
    add_heading_ar(doc, "2.  مصادر البيانات")
    add_paragraph_ar(
        doc,
        "تعتمد بيانات هذا السجل حصراً على الوثائق التالية، وهي الوثائق الثلاث الوحيدة "
        "المعتمدة كمصدر بيانات فعلي لأدلة مؤشرات DC.M.12؛ أما DC.M.12 ذاته فيُستخدم فقط "
        "للتحقق من مطابقة النتيجة النهائية، دون استخدامه كمصدر بيانات خام:",
    )
    for s in rec["sources"]:
        add_bullet_ar(doc, s)
    doc.add_paragraph()

    # ---- 4) جدول البيانات المستخدمة ----
    add_heading_ar(doc, "3.  جدول البيانات المستخدمة")
    add_paragraph_ar(doc, rec["data_table_note"])
    col_widths = [Cm(w) for w in rec["data_table_col_widths_cm"]]
    make_table(
        doc,
        rec["data_table_headers"],
        rec["data_table_rows"],
        col_widths,
        status_col=rec["data_table_status_col"],
    )
    doc.add_paragraph()

    section_num = 4

    if is_calc:
        # ---- 5) تطبيق المعادلة ----
        add_heading_ar(doc, f"{section_num}.  تطبيق المعادلة")
        calc = rec["calc"]
        make_table(
            doc,
            ["البند", "القيمة"],
            [
                (calc["numerator_label"], calc["numerator_value"]),
                (calc["denominator_label"], calc["denominator_value"]),
                ("المعادلة", calc["formula_text"]),
                ("تطبيق المعادلة", calc["application_text"]),
                ("النتيجة النهائية", calc["result"]),
            ],
            [Cm(9.5), Cm(6.5)],
            status_col=[1],
        )
        if rec.get("calc_note"):
            doc.add_paragraph()
            add_paragraph_ar(doc, rec["calc_note"], size=10, italic=True, color=GREY)
        doc.add_paragraph()
        section_num += 1
    else:
        # ---- 5) القيمة المُثبَتة وحدود البيانات ----
        add_heading_ar(doc, f"{section_num}.  القيمة المُثبَتة وحدود البيانات المتاحة")
        add_paragraph_ar(doc, "القيمة الإجمالية كما وردت في المصدر:", bold=True, size=11.5)
        add_paragraph_ar(doc, rec["evidence_reported_value"])
        doc.add_paragraph()
        add_paragraph_ar(doc, "ملاحظة بشأن البيانات التفصيلية غير المتاحة:", bold=True, size=11.5)
        add_paragraph_ar(doc, rec["evidence_unavailable_note"], size=10.5, color=GREY)
        doc.add_paragraph()
        section_num += 1

    # ---- 6) النتيجة والتحقق مقابل DC.M.12 ----
    add_heading_ar(doc, f"{section_num}.  النتيجة والتحقق مقابل DC.M.12")
    make_table(
        doc,
        ["البند", "القيمة"],
        [
            ("القيمة المستهدفة (DC.M.12)", rec["target"]),
            ("القيمة الفعلية المُثبَتة في DC.M.12", rec["dcm12_actual"]),
            ("حالة الإنجاز في DC.M.12", rec["dcm12_status"]),
            ("مرجع التحقق في DC.M.12", rec["dcm12_ref"]),
        ],
        [Cm(6.5), Cm(9.5)],
        status_col=[1],
    )
    doc.add_paragraph()
    add_paragraph_ar(doc, rec["verification_text"])
    doc.add_paragraph()
    section_num += 1

    # ---- 7) الاعتماد ----
    add_heading_ar(doc, f"{section_num}.  الاعتماد")
    add_paragraph_ar(
        doc,
        "تمت مراجعة هذا السجل من مكتب إدارة البيانات، وتم التحقق من مطابقة البيانات "
        f"الواردة فيه لمحتوى {' و'.join(s.split(' — ')[0] for s in rec['sources'])} دون أي "
        "إضافة أو حذف أو تعديل، ومن مطابقة نتيجته النهائية للقيمة الفعلية الموثقة في DC.M.12.",
    )
    doc.add_paragraph()

    approval_table = doc.add_table(rows=2, cols=4)
    approval_table.style = "Table Grid"
    set_table_rtl(approval_table)
    headers = ["صاحب الصلاحية", "المسمى الوظيفي", "التوقيع", "التاريخ"]
    hdr_cells = approval_table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_background(hdr_cells[i], "1F2A44")
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = hdr_cells[i].paragraphs[0]
        rtl(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_rtl_run(r)

    vals = ["________________", "مدير مكتب إدارة البيانات", "________________", "________________"]
    row_cells = approval_table.rows[1].cells
    for i, v in enumerate(vals):
        p = row_cells[i].paragraphs[0]
        rtl(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(v)
        r.font.size = Pt(10.5)
        set_rtl_run(r)

    doc.add_paragraph()
    add_paragraph_ar(
        doc,
        f"ملاحظة: هذا السجل دليل داعم مستقل ضمن حزمة أدلة تصنيف البيانات الخاصة بهيئة الخدمات الحكومية الذكية (SGSA)، أُعدّ لتوثيق بيانات "
        f"وطريقة {'احتساب' if is_calc else 'إثبات'} مؤشر {rec['kpi_code']} دون تعديل أي "
        "محتوى في تقرير DC.M.12 أو الوثائق المصدر.",
        size=9.5, italic=True, color=GREY,
    )

    doc.save(out_path)
    print("Saved:", out_path)


if __name__ == "__main__":
    # نطاق طلب "تعديل تنسيق صارم" (E05/E06/E07 فقط) — لا يُعاد بناء E01-E04 هنا حتى لا
    # تُمسّ ملفاتها القائمة خارج النطاق المطلوب صراحةً. لإعادة بناء الحزمة كاملةً لاحقاً
    # بنفس قالب التنسيق المُحدَّث، بدّل TARGET_IDS إلى None.
    TARGET_IDS = {"E05", "E06", "E07"}
    os.makedirs(OUT_DIR, exist_ok=True)
    for rec in RECORDS:
        if TARGET_IDS is not None and rec["id"] not in TARGET_IDS:
            continue
        build_record(rec)
