# -*- coding: utf-8 -*-
"""XLSX twins of the 7 DC.M.12 KPI docx/pdf records.
See _build_dc_m12_kpi_docx.py / _dc_m12_kpi_data.py for the shared data and rationale.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from _dc_m12_kpi_data import RECORDS, DOC_META_COMMON

OUT_DIR = "06_DC_M12_KPI_Evidence"

NAVY = "1F2A44"
GOLD = "AB8654"
LIGHT_GREY = "F2F2F2"
WHITE = "FFFFFF"
GREEN = "1F6F43"
AMBER = "8A6D00"
GREY_TXT = "595959"
AR_FONT = "Arial"

thin = Side(style="thin", color="D9D9D9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def status_color(v):
    return GREEN if str(v) in ("تثبيت التصنيف", "مُنشور", "محقق") else AMBER


def build_record(rec):
    is_calc = rec["record_type"] == "calculation"
    out_path = os.path.join(OUT_DIR, rec["docx_name"] + ".xlsx")
    n_cols = max(8, len(rec["data_table_headers"]))

    wb = Workbook()
    ws = wb.active
    ws.title = rec["kpi_code"][:31]
    ws.sheet_view.showGridLines = False
    ws.sheet_view.rightToLeft = True

    r = 1

    def merged_text(row, text, size, bold, color, fill=None, height=None, italic=False, span=n_cols):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = Font(name=AR_FONT, size=size, bold=bold, italic=italic, color=color)
        cell.alignment = Alignment(horizontal="right", vertical="center", indent=1, wrap_text=True)
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        if height:
            ws.row_dimensions[row].height = height
        return row + 1

    def section_heading(row, text, span=n_cols):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = Font(name=AR_FONT, size=12.5, bold=True, color=NAVY)
        cell.alignment = Alignment(horizontal="right", vertical="center", indent=1)
        cell.border = Border(bottom=Side(style="medium", color=GOLD))
        ws.row_dimensions[row].height = 22
        return row + 1

    def spacer(row, h=10):
        ws.row_dimensions[row].height = h
        return row + 1

    def meta_row(row, label, value, span=n_cols):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        c1 = ws.cell(row=row, column=1, value=label)
        c1.font = Font(name=AR_FONT, size=10, bold=True, color=GREY_TXT)
        c1.alignment = Alignment(horizontal="right", vertical="center", indent=1)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=span)
        c2 = ws.cell(row=row, column=3, value=value)
        c2.font = Font(name=AR_FONT, size=10, color="262626")
        c2.alignment = Alignment(horizontal="right", vertical="center", indent=1, wrap_text=True)
        return row + 1

    def wrapped_note(row, text, span=n_cols, size=9, italic=True, color="7F7F7F", height=30):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = Font(name=AR_FONT, size=size, italic=italic, color=color)
        cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True, indent=1)
        ws.row_dimensions[row].height = height
        return row + 1

    def data_table(row, headers, rows, status_col_idx=None):
        header_row = row
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=c, value=h)
            cell.font = Font(name=AR_FONT, size=9.5, bold=True, color=WHITE)
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BORDER
        ws.row_dimensions[header_row].height = 30
        row += 1
        for i, row_vals in enumerate(rows):
            row_fill = LIGHT_GREY if i % 2 == 0 else WHITE
            for c, v in enumerate(row_vals, start=1):
                cell = ws.cell(row=row, column=c, value=v)
                cell.border = BORDER
                cell.fill = PatternFill("solid", fgColor=row_fill)
                is_status = status_col_idx is not None and c == status_col_idx + 1
                cell.alignment = Alignment(
                    horizontal="center" if (c == 1 or is_status) else "right",
                    vertical="center", wrap_text=True,
                    indent=0 if (c == 1 or is_status) else 1,
                )
                if c == 1:
                    cell.font = Font(name=AR_FONT, size=9.5, bold=True, color=NAVY)
                elif is_status:
                    cell.font = Font(name=AR_FONT, size=9.5, bold=True, color=status_color(v))
                else:
                    cell.font = Font(name=AR_FONT, size=9, color="262626")
            ws.row_dimensions[row].height = 20
            row += 1
        return row

    def kv_table(row, rows, value_bold_color=None):
        header_row = row
        for c, h in enumerate(["البند", "القيمة"], start=1):
            span_end = n_cols if c == 2 else 1
            if c == 2:
                ws.merge_cells(start_row=header_row, start_column=2, end_row=header_row, end_column=n_cols)
            cell = ws.cell(row=header_row, column=1 if c == 1 else 2, value=h)
            cell.font = Font(name=AR_FONT, size=9.5, bold=True, color=WHITE)
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER
        ws.row_dimensions[header_row].height = 20
        row += 1
        for i, (label, value) in enumerate(rows):
            row_fill = LIGHT_GREY if i % 2 == 0 else WHITE
            c1 = ws.cell(row=row, column=1, value=label)
            c1.font = Font(name=AR_FONT, size=9.5, bold=True, color=NAVY)
            c1.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True, indent=1)
            c1.fill = PatternFill("solid", fgColor=row_fill)
            c1.border = BORDER
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=n_cols)
            c2 = ws.cell(row=row, column=2, value=value)
            c2.font = Font(name=AR_FONT, size=9.5, bold=True, color=value_bold_color or "262626")
            c2.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True, indent=1)
            c2.fill = PatternFill("solid", fgColor=row_fill)
            for cc in range(2, n_cols + 1):
                ws.cell(row=row, column=cc).border = BORDER
            ws.row_dimensions[row].height = 22
            row += 1
        return row

    # ============ Title banner ============
    r = merged_text(r, "هيئة الخدمات الحكومية الذكية (SGSA)", 18, True, WHITE, fill=NAVY, height=32)
    r = merged_text(r, "منصة ذكية لإدارة وقياس الامتثال لمتطلبات المكتب الوطني لإدارة البيانات (NDMO)", 10.5, True, WHITE, fill=NAVY, height=20)
    r = spacer(r, 6)
    record_type_ar = "سجل احتساب" if is_calc else "سجل إثبات"
    r = merged_text(r, f"{record_type_ar} مؤشر «{rec['kpi_name_ar']}» ({rec['kpi_code']})", 13.5, True, NAVY, height=26)
    suffix_en = "KPI Calculation Record" if is_calc else "KPI Evidence Record"
    r = merged_text(r, f"دليل داعم — {rec['doc_code']}  |  {suffix_en}", 10.5, True, GOLD, italic=True, height=18)
    r = spacer(r, 10)

    # ============ 1) بيانات الوثيقة ============
    r = section_heading(r, "1.  بيانات الوثيقة")
    r = spacer(r, 4)
    r = meta_row(r, "الجهة", DOC_META_COMMON["entity"])
    r = meta_row(r, "المنصة", DOC_META_COMMON["platform"])
    r = meta_row(r, "الإطار المرجعي", DOC_META_COMMON["framework"])
    r = meta_row(r, "رمز الدليل", rec["doc_code"])
    r = meta_row(r, "المتطلب المرتبط", DOC_META_COMMON["related_requirement"])
    r = meta_row(r, "المؤشر المرتبط", f"{rec['kpi_code']} — {rec['kpi_name_ar']}")
    r = meta_row(r, "نوع السجل", "سجل احتساب (Calculation Record)" if is_calc else "سجل إثبات (Evidence Record)")
    r = meta_row(r, "الإصدار", DOC_META_COMMON["version"])
    r = meta_row(r, "التاريخ", DOC_META_COMMON["date"])
    r = meta_row(r, "حالة الوثيقة", DOC_META_COMMON["status"])
    r = spacer(r, 14)

    # ============ 2) نطاق القياس ============
    r = section_heading(r, "2.  نطاق القياس")
    r = spacer(r, 4)
    r = wrapped_note(r, rec["scope"], size=10, italic=False, color="262626", height=42)
    r = spacer(r, 14)

    # ============ 3) مصادر البيانات ============
    r = section_heading(r, "3.  مصادر البيانات")
    r = spacer(r, 4)
    r = wrapped_note(
        r,
        "تعتمد بيانات هذا السجل حصراً على الوثائق التالية، وهي الوثائق الثلاث الوحيدة "
        "المعتمدة كمصدر بيانات فعلي لأدلة مؤشرات DC.M.12؛ أما DC.M.12 ذاته فيُستخدم فقط "
        "للتحقق من مطابقة النتيجة النهائية، دون استخدامه كمصدر بيانات خام:",
        size=10, italic=False, color="262626", height=34,
    )
    for s in rec["sources"]:
        r = wrapped_note(r, "•  " + s, size=9.5, italic=False, color="262626", height=30)
    r = spacer(r, 14)

    # ============ 4) جدول البيانات المستخدمة ============
    r = section_heading(r, "4.  جدول البيانات المستخدمة")
    r = spacer(r, 4)
    r = wrapped_note(r, rec["data_table_note"], size=9.5, italic=False, color="262626", height=20)
    status_idx = rec["data_table_status_col"][0] if rec["data_table_status_col"] else None
    r = data_table(r, rec["data_table_headers"], rec["data_table_rows"], status_col_idx=status_idx)
    r = spacer(r, 16)

    section_num = 5
    if is_calc:
        r = section_heading(r, f"{section_num}.  تطبيق المعادلة")
        r = spacer(r, 4)
        calc = rec["calc"]
        r = kv_table(r, [
            (calc["numerator_label"], calc["numerator_value"]),
            (calc["denominator_label"], calc["denominator_value"]),
            ("المعادلة", calc["formula_text"]),
            ("تطبيق المعادلة", calc["application_text"]),
            ("النتيجة النهائية", calc["result"]),
        ], value_bold_color=GREEN)
        if rec.get("calc_note"):
            r = spacer(r, 6)
            r = wrapped_note(r, rec["calc_note"], height=34)
        r = spacer(r, 16)
        section_num += 1
    else:
        r = section_heading(r, f"{section_num}.  القيمة المُثبَتة وحدود البيانات المتاحة")
        r = spacer(r, 4)
        r = wrapped_note(r, "القيمة الإجمالية كما وردت في المصدر:", size=10, italic=False, color=NAVY, height=18)
        r = wrapped_note(r, rec["evidence_reported_value"], size=10, italic=False, color="262626", height=34)
        r = spacer(r, 6)
        r = wrapped_note(r, "ملاحظة بشأن البيانات التفصيلية غير المتاحة:", size=10, italic=False, color=NAVY, height=18)
        r = wrapped_note(r, rec["evidence_unavailable_note"], height=60)
        r = spacer(r, 16)
        section_num += 1

    # ============ النتيجة والتحقق مقابل DC.M.12 ============
    r = section_heading(r, f"{section_num}.  النتيجة والتحقق مقابل DC.M.12")
    r = spacer(r, 4)
    r = kv_table(r, [
        ("القيمة المستهدفة (DC.M.12)", rec["target"]),
        ("القيمة الفعلية المُثبَتة في DC.M.12", rec["dcm12_actual"]),
        ("حالة الإنجاز في DC.M.12", rec["dcm12_status"]),
        ("مرجع التحقق في DC.M.12", rec["dcm12_ref"]),
    ], value_bold_color=GREEN)
    r = spacer(r, 6)
    r = wrapped_note(r, rec["verification_text"], size=10, italic=False, color="262626", height=34)
    r = spacer(r, 16)
    section_num += 1

    # ============ الاعتماد ============
    r = section_heading(r, f"{section_num}.  الاعتماد")
    r = spacer(r, 4)
    sources_short = " و".join(s.split(" — ")[0] for s in rec["sources"])
    r = wrapped_note(
        r,
        "تمت مراجعة هذا السجل من مكتب إدارة البيانات، وتم التحقق من مطابقة البيانات "
        f"الواردة فيه لمحتوى {sources_short} دون أي إضافة أو حذف أو تعديل، ومن مطابقة "
        "نتيجته النهائية للقيمة الفعلية الموثقة في DC.M.12.",
        size=10, italic=False, color="262626", height=34,
    )
    r = spacer(r, 6)

    approval_headers = ["صاحب الصلاحية", "المسمى الوظيفي", "التوقيع", "التاريخ"]
    quarter = max(2, n_cols // 4)
    col_spans = []
    start = 1
    for i in range(4):
        end = start + quarter - 1 if i < 3 else n_cols
        col_spans.append((start, min(end, n_cols)))
        start = end + 1
    approval_header_row = r
    for (sc, ec), h in zip(col_spans, approval_headers):
        ec = max(ec, sc)
        ws.merge_cells(start_row=approval_header_row, start_column=sc, end_row=approval_header_row, end_column=ec)
        cell = ws.cell(row=approval_header_row, column=sc, value=h)
        cell.font = Font(name=AR_FONT, size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        for cc in range(sc, ec + 1):
            ws.cell(row=approval_header_row, column=cc).border = BORDER
    ws.row_dimensions[approval_header_row].height = 20
    r += 1

    approval_values = ["____________", "مدير مكتب إدارة البيانات", "____________", "____________"]
    for (sc, ec), v in zip(col_spans, approval_values):
        ec = max(ec, sc)
        ws.merge_cells(start_row=r, start_column=sc, end_row=r, end_column=ec)
        cell = ws.cell(row=r, column=sc, value=v)
        cell.font = Font(name=AR_FONT, size=9.5, color="262626")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        for cc in range(sc, ec + 1):
            ws.cell(row=r, column=cc).border = BORDER
    ws.row_dimensions[r].height = 26
    r += 1

    r = spacer(r, 14)
    r = wrapped_note(
        r,
        f"ملاحظة: هذا السجل دليل داعم مستقل ضمن حزمة أدلة تصنيف البيانات الخاصة بهيئة الخدمات الحكومية الذكية (SGSA)، أُعدّ لتوثيق بيانات "
        f"وطريقة {'احتساب' if is_calc else 'إثبات'} مؤشر {rec['kpi_code']} دون تعديل أي "
        "محتوى في تقرير DC.M.12 أو الوثائق المصدر.",
        height=26,
    )

    # ============ Column widths ============
    n_data_cols = len(rec["data_table_headers"])
    widths = [14] + [22] * (n_data_cols - 1) if n_data_cols > 1 else [30]
    while len(widths) < n_cols:
        widths.append(14)
    for i, w in enumerate(widths[:n_cols], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(out_path)
    print("Saved:", out_path)


if __name__ == "__main__":
    os.makedirs(OUT_DIR, exist_ok=True)
    for rec in RECORDS:
        build_record(rec)
