# -*- coding: utf-8 -*-
"""Builds a NEW, standalone draft Evidence Artifact for DC.M.2 within the Wathiq/SGSA
portfolio context — independent from any existing evidence_repository file.

Draft status: for review before any later integration. Does not modify or
replace any existing file. Arabic / RTL, same visual identity as the rest of
wathiq_portfolio_package (Wathiq = platform, SGSA = entity under review, NDMO = framework).
"""
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_PATH = "05_DC_M2_Draft_Evidence/DC_M2_Implementation_Status_Report.docx"

NAVY = RGBColor(0x1F, 0x2A, 0x44)
GOLD = RGBColor(0xAB, 0x8D, 0x54)
GREY = RGBColor(0x59, 0x59, 0x59)
GREEN = RGBColor(0x1F, 0x6F, 0x43)
AMBER = RGBColor(0x8A, 0x6D, 0x00)
AR_FONT = "Arial"


def rtl(p):
    b = OxmlElement("w:bidi")
    b.set(qn("w:val"), "1")
    p._p.get_or_add_pPr().append(b)


def set_rtl_run(run):
    pr = run._element.get_or_add_rPr()
    pr.get_or_add_rFonts().set(qn("w:cs"), AR_FONT)
    e = OxmlElement("w:rtl")
    e.set(qn("w:val"), "1")
    pr.append(e)


def set_table_rtl(table):
    tbl_pr = table._tbl.tblPr
    bv = OxmlElement("w:bidiVisual")
    bv.set(qn("w:val"), "1")
    tbl_pr.append(bv)


def set_cell_background(cell, color_hex):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color_hex)
    cell._tc.get_or_add_tcPr().append(shd)


def add_meta_row(table, label, value):
    row = table.add_row()
    p1 = row.cells[1].paragraphs[0]
    rtl(p1)
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p1.add_run(label)
    r1.font.bold = True
    r1.font.size = Pt(9)
    r1.font.color.rgb = GREY
    set_rtl_run(r1)

    p2 = row.cells[0].paragraphs[0]
    rtl(p2)
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = p2.add_run(value)
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
    set_rtl_run(r2)


def add_paragraph_ar(doc, text, size=9.5, bold=False, italic=False, color=RGBColor(0x26, 0x26, 0x26)):
    p = doc.add_paragraph()
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    set_rtl_run(r)
    return p


def add_heading_ar(doc, text, size=13):
    p = doc.add_paragraph()
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = NAVY
    set_rtl_run(r)
    return p


def add_bullet_ar(doc, text, size=9.5):
    p = doc.add_paragraph()
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("•  " + text)
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
    set_rtl_run(r)
    return p


def make_table(doc, headers, rows, col_widths, status_col=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_rtl(table)
    for i, w in enumerate(col_widths):
        table.columns[i].width = w

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        set_cell_background(hdr_cells[i], "1F2A44")
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = hdr_cells[i].paragraphs[0]
        rtl(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_rtl_run(r)

    for idx, row_vals in enumerate(rows):
        row = table.add_row()
        for c, v in enumerate(row_vals):
            cell = row.cells[c]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if idx % 2 == 0:
                set_cell_background(cell, "F2F2F2")
            p = cell.paragraphs[0]
            rtl(p)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c in (status_col or []) else WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(str(v))
            run.font.size = Pt(8.5)
            set_rtl_run(run)
            if c == 0:
                run.font.bold = True
                run.font.color.rgb = NAVY
            elif status_col and c in status_col:
                run.font.bold = True
                run.font.color.rgb = GREEN if v == "مكتمل" or v == "متوفر" else AMBER
            else:
                run.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
    return table


doc = docx.Document()
section = doc.sections[0]
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)
section.top_margin = Cm(1.5)
section.bottom_margin = Cm(1.5)

style = doc.styles["Normal"]
style.font.name = AR_FONT
style.font.size = Pt(10)
style.element.rPr.rFonts.set(qn("w:cs"), AR_FONT)

# ============ 1) صفحة الغلاف ============
banner = doc.add_table(rows=1, cols=1)
banner.autofit = True
set_table_rtl(banner)
cell = banner.rows[0].cells[0]
set_cell_background(cell, "1F2A44")
cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
p = cell.paragraphs[0]
rtl(p)
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = p.add_run("تقرير حالة تنفيذ خطة تصنيف البيانات")
run.font.size = Pt(18)
run.font.bold = True
run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
set_rtl_run(run)
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(6)

doc.add_paragraph()
add_paragraph_ar(doc, "منصة ذكية لإدارة وقياس الامتثال لمتطلبات المكتب الوطني لإدارة البيانات (NDMO)",
                  size=10.5, bold=True, color=GOLD)
add_heading_ar(doc, "تقرير حالة تنفيذ خطة تصنيف البيانات", size=15)
add_paragraph_ar(doc, "مسودة دليل مستقل للمراجعة — DC.M.2", size=11, italic=True, color=GOLD)

doc.add_paragraph()

meta_table = doc.add_table(rows=0, cols=2)
meta_table.autofit = True
set_table_rtl(meta_table)
meta_table.columns[0].width = Cm(11)
meta_table.columns[1].width = Cm(5)
add_meta_row(meta_table, "الجهة", "هيئة الخدمات الحكومية الذكية (SGSA)")
add_meta_row(meta_table, "القسم", "مكتب إدارة البيانات")
add_meta_row(meta_table, "رمز الدليل", "DC.M.2")
add_meta_row(meta_table, "اسم الدليل", "تقرير حالة تنفيذ خطة تصنيف البيانات")
add_meta_row(meta_table, "الإصدار", "0.1 (مسودة)")
add_meta_row(meta_table, "التاريخ", "يونيو 2026")
add_meta_row(meta_table, "حالة الوثيقة", "مسودة للمراجعة — لم تُعتمد بعد")
add_meta_row(meta_table, "الإطار المرجعي", "إطار حوكمة البيانات الوطني (NDMO)")
add_meta_row(meta_table, "مرجع الخطة المعتمدة", "DC.C.1.1")

doc.add_page_break()

# ============ 2) المقدمة ============
add_heading_ar(doc, "1.  المقدمة")
add_paragraph_ar(
    doc,
    "يُعدّ متطلب DC.M.2 ضمن سؤال النضج DC.MQ.1 مطلباً أساسياً في المرحلة الأولى (البناء) من "
    "أطوار تصنيف البيانات، إذ ينص على أن: «تعمل الجهة على تنفيذ الخطة المعرفة وخارطة طريق "
    "تصنيف البيانات على جميع المجموعات المعتمدة من البيانات والسجلات». يوثّق هذا التقرير "
    "حالة تنفيذ هذه الخطة فعلياً داخل هيئة الخدمات الحكومية الذكية (SGSA)، ويُعدّ الدليل "
    "الرسمي المُستخدَم لإثبات استيفاء هذا المتطلب."
)
add_paragraph_ar(
    doc,
    "يستعرض التقرير خارطة الطريق المعتمدة ونطاق تنفيذها، ونطاق التغطية الفعلي عبر مجموعات "
    "البيانات والسجلات والإدارات المالكة، وحالة كل نشاط من أنشطة الخطة، والمبادرات المنفذة "
    "فعلياً، ومؤشرات التقدم المحققة، إضافة إلى الأدلة الداعمة لكل ما سبق."
)

doc.add_paragraph()

# ============ 3) الخطة المعتمدة وخارطة الطريق ============
add_heading_ar(doc, "2.  الخطة المعتمدة وخارطة الطريق")
add_paragraph_ar(doc, "أهداف الخطة:", bold=True, size=10)
for b in [
    "حصر جميع مجموعات البيانات والسجلات المعتمدة عبر الإدارات.",
    "تحديد وتوثيق ملاك البيانات ومسؤولياتهم رسمياً.",
    "رفع مستوى الوعي المؤسسي بمفاهيم تصنيف البيانات ومتطلباته.",
    "إعداد إجراءات تشغيلية موحَّدة لتصنيف البيانات تمهيداً لاعتمادها.",
]:
    add_bullet_ar(doc, b)

doc.add_paragraph()
add_paragraph_ar(doc, "نطاق التنفيذ:", bold=True, size=10)
add_paragraph_ar(
    doc,
    "يشمل نطاق التنفيذ جميع مجموعات البيانات والسجلات المعتمدة ضمن الجرد المؤسسي، عبر خمس "
    "إدارات رئيسية مشمولة بالخطة (مكتب إدارة البيانات، الموارد البشرية، الإدارة المالية، "
    "تقنية المعلومات، الخدمات الرقمية)."
)

doc.add_paragraph()
add_paragraph_ar(doc, "الفترة الزمنية:", bold=True, size=10)
add_paragraph_ar(
    doc,
    "تغطي هذه المرحلة (مرحلة البناء — Level 1) الفترة من الربع الأول 2026 وحتى تاريخ إصدار "
    "هذا التقرير في يونيو 2026، تمهيداً للانتقال إلى مرحلة التعريف (Level 2)."
)

doc.add_paragraph()

# ============ 4) نطاق التغطية ============
add_heading_ar(doc, "3.  نطاق التغطية")
add_paragraph_ar(doc, "يوضّح الجدول التالي نطاق التغطية الفعلي لخطة تصنيف البيانات عبر الإدارات المالكة:")
make_table(
    doc,
    ["مجموعات البيانات المشمولة", "السجلات المشمولة", "الإدارة المالكة", "حالة التنفيذ"],
    [
        ("بيانات الموظفين والملفات الوظيفية", "سجلات العقود والاتفاقيات", "إدارة الموارد البشرية", "مكتمل"),
        ("بيانات المدفوعات والمعاملات المالية", "سجلات التدقيق الداخلي والرقابة", "الإدارة المالية", "مكتمل"),
        ("بيانات الأصول والبنية التحتية التقنية", "سجلات التغييرات والتحديثات التقنية", "إدارة تقنية المعلومات", "مكتمل"),
        ("بيانات المستخدمين والمستفيدين", "سجلات الشكاوى والمقترحات", "إدارة الخدمات الرقمية", "قيد التنفيذ"),
        ("بيانات التقارير والإحصاءات المنشورة", "سجلات التقارير الإدارية الدورية", "مكتب إدارة البيانات", "مكتمل"),
    ],
    [Cm(4.8), Cm(4.8), Cm(4.0), Cm(2.5)],
    status_col=[3],
)

doc.add_paragraph()

# ============ 5) حالة تنفيذ الخطة ============
add_heading_ar(doc, "4.  حالة تنفيذ الخطة")
make_table(
    doc,
    ["النشاط", "المسؤول", "الحالة", "نسبة الإنجاز", "الملاحظات"],
    [
        ("حصر مجموعات البيانات والسجلات", "مكتب إدارة البيانات", "مكتمل", "100%",
         "تم حصر وتوثيق جميع المجموعات الرئيسية دون تحديات جوهرية"),
        ("تحديد وتوثيق ملاك البيانات", "مكتب إدارة البيانات", "مكتمل", "100%",
         "صدر قرار رسمي بتحديد ملاك البيانات في جميع الإدارات"),
        ("تنفيذ برنامج التوعية والتدريب", "إدارة الموارد البشرية", "قيد التنفيذ", "90%",
         "تم تنفيذ معظم الورش المخططة؛ تبقّت جلسة واحدة لإدارة الخدمات الرقمية"),
        ("إعداد إجراءات التصنيف التشغيلية", "مكتب إدارة البيانات", "قيد التنفيذ", "80%",
         "الإجراءات في مرحلة المراجعة النهائية تمهيداً للاعتماد الرسمي"),
        ("عقد اجتماعات المتابعة الدورية", "مكتب إدارة البيانات", "مكتمل", "100%",
         "عُقدت الاجتماعات المقررة وتم توثيقها بمحاضر رسمية"),
    ],
    [Cm(3.6), Cm(3.0), Cm(2.2), Cm(1.8), Cm(5.5)],
    status_col=[2],
)

doc.add_paragraph()

# ============ 6) المبادرات والأنشطة المنفذة ============
add_heading_ar(doc, "5.  المبادرات والأنشطة المنفذة")
add_paragraph_ar(doc, "فيما يلي تفاصيل الإنجازات الفعلية لكل نشاط من أنشطة الخطة:")
add_bullet_ar(doc, "حصر مجموعات البيانات: أُجري جرد شامل لمجموعات البيانات والسجلات الرئيسية عبر الإدارات الخمس، وأُعدَّت قائمة موثَّقة تشمل اسم المجموعة والإدارة المالكة ودرجة أهميتها التشغيلية.")
add_bullet_ar(doc, "توثيق ملاك البيانات: صدر قرار رسمي بتحديد ملاك البيانات، وتضمَّن السجل المعتمَد اسم كل مالك ومسؤولياته والبيانات التابعة لنطاق إدارته.")
add_bullet_ar(doc, "ورش التوعية الداخلية: نُفِّذت ورش توعية استهدفت منسوبي الإدارات المشمولة حول أهمية تصنيف البيانات ومستوياته وآليات التطبيق العملي.")
add_bullet_ar(doc, "مسودة إجراءات التصنيف: أُعِدَّت مسودة الإجراءات التشغيلية الموحَّدة لتصنيف البيانات، وهي قيد المراجعة النهائية قبل رفعها للاعتماد.")
add_bullet_ar(doc, "اجتماعات المتابعة: عُقدت اجتماعات دورية بمشاركة مكتب إدارة البيانات وممثلي الإدارات المعنية لمتابعة سير التنفيذ وتوثيق القرارات.")

doc.add_paragraph()

# ============ 7) مؤشرات التقدم ============
add_heading_ar(doc, "6.  مؤشرات التقدم")
make_table(
    doc,
    ["المؤشر", "القيمة"],
    [
        ("نسبة إنجاز الخطة الإجمالية", "94%"),
        ("عدد الأنشطة المكتملة", "3 من 5"),
        ("عدد الإدارات المشمولة", "5 إدارات"),
    ],
    [Cm(10.0), Cm(6.0)],
    status_col=[1],
)

doc.add_paragraph()

# ============ 8) الأدلة الداعمة ============
add_heading_ar(doc, "7.  الأدلة الداعمة")
make_table(
    doc,
    ["نوع الدليل", "الوصف", "الحالة"],
    [
        ("تقرير حصر مجموعات البيانات", "وثيقة تحتوي القائمة الشاملة لمجموعات البيانات المحصورة", "متوفر"),
        ("سجل ملاك البيانات", "سجل رسمي بأسماء ملاك البيانات ومسؤولياتهم لكل إدارة", "متوفر"),
        ("مواد التدريب والتوعية", "العروض والمواد المستخدمة في ورش التوعية المنفَّذة", "متوفر"),
        ("محاضر اجتماعات المتابعة", "محاضر موثَّقة لاجتماعات متابعة تنفيذ الخطة", "متوفر"),
        ("مسودة إجراءات التصنيف", "المسودة الحالية للإجراءات التشغيلية لتصنيف البيانات", "قيد المراجعة"),
    ],
    [Cm(4.5), Cm(8.5), Cm(3.0)],
    status_col=[2],
)

doc.add_paragraph()
doc.add_paragraph()

# ============ 9) الاعتماد ============
add_heading_ar(doc, "8.  الاعتماد")
add_paragraph_ar(
    doc,
    "هذه الوثيقة مسودة قيد المراجعة، ولم تُعرَض بعد على صاحب الصلاحية للاعتماد الرسمي. "
    "يُستكمَل الاعتماد أدناه عند إقرار المسودة ودمجها ضمن مستودع الأدلة."
)
doc.add_paragraph()

approval_table = doc.add_table(rows=2, cols=4)
approval_table.style = "Table Grid"
set_table_rtl(approval_table)
headers = ["صاحب الصلاحية", "المسمى الوظيفي", "التوقيع", "التاريخ"]
hdr_cells = approval_table.rows[0].cells
for i, h in enumerate(headers):
    set_cell_background(hdr_cells[i], "1F2A44")
    hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = hdr_cells[i].paragraphs[0]
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(h)
    r.font.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    set_rtl_run(r)

vals = ["...............", "مدير عام الهيئة", "____________", "بانتظار الاعتماد"]
row_cells = approval_table.rows[1].cells
for i, v in enumerate(vals):
    p = row_cells[i].paragraphs[0]
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(v)
    r.font.size = Pt(9)
    set_rtl_run(r)

doc.add_paragraph()
add_paragraph_ar(
    doc,
    "ملاحظة: هذه الوثيقة مسودة دليل مستقلة ضمن حزمة أدلة تصنيف البيانات الخاصة بهيئة الخدمات الحكومية الذكية (SGSA) لأغراض المراجعة، ولم تُدمَج "
    "بعد ضمن مستودع الأدلة الرسمي.",
    size=8, italic=True, color=GREY,
)

doc.save(OUT_PATH)
print("Saved:", OUT_PATH)
