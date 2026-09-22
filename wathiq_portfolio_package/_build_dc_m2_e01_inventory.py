# -*- coding: utf-8 -*-
"""Builds a NEW, standalone evidence artifact: DC.M.2-E01 Data Inventory Register (Excel).

Independent file — does not modify or replace DC_M2_Implementation_Status_Report.docx/.pdf
or any other existing file. Wathiq/SGSA identity, RTL Arabic, Audit-Evidence-ready formatting.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT_PATH = "05_DC_M2_Draft_Evidence/DC.M.2-E01_Data_Inventory_Register.xlsx"

NAVY = "1F2A44"
GOLD = "AB8654"
LIGHT_GREY = "F2F2F2"
WHITE = "FFFFFF"
GREEN = "1F6F43"
AR_FONT = "Arial"

INVENTORY_ROWS = [
    (1, "بيانات الموظفين", "ملفات الموظفين والعقود الوظيفية", "إدارة الموارد البشرية", "نظام الموارد البشرية", "مكتمل"),
    (2, "بيانات المعاملات المالية", "الفواتير وأوامر الصرف وسجلات الدفع", "الإدارة المالية", "النظام المالي", "مكتمل"),
    (3, "بيانات الأصول والبنية التحتية التقنية", "سجلات الأجهزة والتغييرات التقنية", "إدارة تقنية المعلومات", "نظام إدارة الأصول", "مكتمل"),
    (4, "بيانات المستفيدين والخدمات", "طلبات الخدمات والشكاوى والبلاغات", "إدارة الخدمات الرقمية", "منصة الخدمات الإلكترونية", "مكتمل"),
    (5, "بيانات التقارير والإحصاءات", "التقارير التشغيلية والدورية", "مكتب إدارة البيانات", "مستودع التقارير", "مكتمل"),
]

wb = Workbook()
ws = wb.active
ws.title = "سجل حصر البيانات"
ws.sheet_view.showGridLines = False
ws.sheet_view.rightToLeft = True

thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

r = 1


def merged_text(row, text, size, bold, color, fill=None, height=None, italic=False, span=6):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name=AR_FONT, size=size, bold=bold, italic=italic, color=color)
    cell.alignment = Alignment(horizontal="right", vertical="center", indent=1, wrap_text=True)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if height:
        ws.row_dimensions[row].height = height
    return row + 1


def section_heading(row, text, span=6):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name=AR_FONT, size=12.5, bold=True, color=NAVY)
    cell.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    cell.border = Border(bottom=Side(style="medium", color=GOLD))
    ws.row_dimensions[row].height = 22
    return row + 1


def spacer(row, h=10):
    ws.row_dimensions[row].height = h
    return row + 1


def meta_row(row, label, value):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    c1 = ws.cell(row=row, column=1, value=label)
    c1.font = Font(name=AR_FONT, size=10, bold=True, color="595959")
    c1.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=6)
    c2 = ws.cell(row=row, column=3, value=value)
    c2.font = Font(name=AR_FONT, size=10, color="262626")
    c2.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    return row + 1


# ============ Title banner ============
r = merged_text(r, "سجل حصر مجموعات البيانات والسجلات", 18, True, WHITE, fill=NAVY, height=32)
r = merged_text(r, "منصة ذكية لإدارة وقياس الامتثال لمتطلبات المكتب الوطني لإدارة البيانات (NDMO)", 10.5, True, WHITE, fill=NAVY, height=20)
r = spacer(r, 6)
r = merged_text(r, "سجل حصر مجموعات البيانات والسجلات", 14, True, NAVY, height=24)
r = merged_text(r, "دليل داعم — DC.M.2-E01", 11, True, GOLD, italic=True, height=18)
r = spacer(r, 10)

# ============ بيانات الوثيقة ============
r = section_heading(r, "بيانات الوثيقة")
r = spacer(r, 4)
r = meta_row(r, "الجهة", "هيئة الخدمات الحكومية الذكية (SGSA)")
r = meta_row(r, "القسم", "مكتب إدارة البيانات")
r = meta_row(r, "المنصة", "")
r = meta_row(r, "الإطار المرجعي", "إطار حوكمة البيانات الوطني (NDMO)")
r = meta_row(r, "رمز الدليل", "DC.M.2-E01")
r = meta_row(r, "المتطلب المرتبط", "DC.M.2 — تنفيذ خطة تصنيف البيانات وخارطة الطريق")
r = meta_row(r, "اسم الدليل", "سجل حصر مجموعات البيانات والسجلات")
r = meta_row(r, "مالك الوثيقة", "مكتب إدارة البيانات (Data Management Office)")
r = meta_row(r, "الإصدار", "1.0")
r = meta_row(r, "التاريخ", "يونيو 2026")
r = meta_row(r, "حالة الوثيقة", "جاهزة للمراجعة")
r = spacer(r, 16)

# ============ Data Inventory Register table ============
r = section_heading(r, "سجل حصر مجموعات البيانات (Data Inventory Register)")
r = spacer(r, 4)

headers = ["الرقم", "مجموعة البيانات", "السجلات المرتبطة", "الإدارة المالكة", "النظام المصدر", "حالة الحصر"]
header_row = r
for c, h in enumerate(headers, start=1):
    cell = ws.cell(row=header_row, column=c, value=h)
    cell.font = Font(name=AR_FONT, size=10.5, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border
ws.row_dimensions[header_row].height = 22
r += 1

for i, row_vals in enumerate(INVENTORY_ROWS):
    row_fill = LIGHT_GREY if i % 2 == 0 else WHITE
    for c, v in enumerate(row_vals, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.border = border
        cell.fill = PatternFill("solid", fgColor=row_fill)
        cell.alignment = Alignment(
            horizontal="center" if c in (1, 6) else "right", vertical="center", wrap_text=True,
            indent=0 if c in (1, 6) else 1,
        )
        if c == 1:
            cell.font = Font(name=AR_FONT, size=10, bold=True, color=NAVY)
        elif c == 2:
            cell.font = Font(name=AR_FONT, size=10, bold=True, color=NAVY)
        elif c == 6:
            cell.font = Font(name=AR_FONT, size=10, bold=True, color=GREEN)
        else:
            cell.font = Font(name=AR_FONT, size=10, color="262626")
    ws.row_dimensions[r].height = 30
    r += 1

r = spacer(r, 20)

# ============ الاعتماد ============
r = section_heading(r, "الاعتماد")
r = spacer(r, 6)

approval_headers = ["المسمى الوظيفي", "الاسم", "التوقيع", "التاريخ"]
approval_header_row = r
col_spans = [(1, 1), (2, 3), (4, 5), (6, 6)]
for (sc, ec), h in zip(col_spans, approval_headers):
    ws.merge_cells(start_row=approval_header_row, start_column=sc, end_row=approval_header_row, end_column=ec)
    cell = ws.cell(row=approval_header_row, column=sc, value=h)
    cell.font = Font(name=AR_FONT, size=10, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    for cc in range(sc, ec + 1):
        ws.cell(row=approval_header_row, column=cc).border = border
ws.row_dimensions[approval_header_row].height = 20
r += 1

approval_values = ["مدير عام الهيئة", "...............", "____________", "________"]
for (sc, ec), v in zip(col_spans, approval_values):
    ws.merge_cells(start_row=r, start_column=sc, end_row=r, end_column=ec)
    cell = ws.cell(row=r, column=sc, value=v)
    cell.font = Font(name=AR_FONT, size=9.5, color="262626")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    for cc in range(sc, ec + 1):
        ws.cell(row=r, column=cc).border = border
ws.row_dimensions[r].height = 26
r += 1

r = spacer(r, 14)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
note = ws.cell(
    row=r, column=1,
    value="ملاحظة: هذا الدليل جزء من حزمة أدلة تصنيف البيانات الخاصة بهيئة الخدمات الحكومية الذكية (SGSA) الداعمة لمتطلب DC.M.2، ويُستخدَم كمرفق تدقيق (Audit Evidence) مستقل.",
)
note.font = Font(name=AR_FONT, size=8, italic=True, color="7F7F7F")
note.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True, indent=1)
ws.row_dimensions[r].height = 26

# ============ Column widths ============
widths = [8, 26, 30, 22, 22, 14]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

wb.save(OUT_PATH)
print("Saved:", OUT_PATH)
