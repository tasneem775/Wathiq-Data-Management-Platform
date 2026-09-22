# -*- coding: utf-8 -*-
"""Updates DC.M.2-E04-A Awareness Workshop Attendance Register (Word) in place.
Independent file — does not modify E04, the main report, E01/E02/E03, or any
other file. Uses 47 generic illustrative Arabic full names (not linked to real
individuals or public figures) instead of numbered placeholders.
Total attendance strictly 47 (16 + 16 + 15), across exactly 3 workshops.
"""
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_PATH = "05_DC_M2_Draft_Evidence/DC.M.2-E04-A_Awareness_Workshop_Attendance_Register.docx"

NAVY = RGBColor(0x1F, 0x2A, 0x44)
GOLD = RGBColor(0xAB, 0x8D, 0x54)
GREY = RGBColor(0x59, 0x59, 0x59)
AR_FONT = "Arial"

DEPARTMENTS = [
    "مكتب إدارة البيانات",
    "إدارة الموارد البشرية",
    "الإدارة المالية",
    "إدارة تقنية المعلومات",
    "إدارة الخدمات الرقمية",
]


def dept_for(i):
    return DEPARTMENTS[i % len(DEPARTMENTS)]


# 47 unique generic Arabic full names (common first name + common family name).
# Not linked to any real individual or public figure — illustrative sample data only.
NAMES = [
    "أحمد السالم", "سارة العتيبي", "خالد القحطاني", "نورة الشهري", "محمد الدوسري",
    "منى الحربي", "عبدالله الغامدي", "فاطمة الزهراني", "فهد المطيري", "هند العنزي",
    "سعود الشمري", "ريم القرني", "ماجد البقمي", "لمياء السبيعي", "طارق الرشيدي",
    "أمل الجهني", "بندر العمري", "عائشة الشهراني", "ناصر الخالدي", "جواهر اليامي",
    "سلطان الفيفي", "لطيفة الثقفي", "عمر السلمي", "مها المالكي", "يوسف الحازمي",
    "غادة العتيبي", "وليد القحطاني", "أروى الدوسري", "تركي الحربي", "دلال الغامدي",
    "فيصل الزهراني", "شيخة المطيري", "عبدالعزيز العنزي", "بشاير الشمري", "راشد القرني",
    "رهف البقمي", "سامي السبيعي", "لينا الرشيدي", "حمد الجهني", "حصة العمري",
    "إبراهيم الشهراني", "سارة الخالدي", "عبدالرحمن اليامي", "نورة الفيفي", "مشعل الثقفي",
    "منى السلمي", "سعد المالكي",
]
assert len(NAMES) == 47
assert len(set(NAMES)) == 47


WORKSHOPS = [
    {
        "title": "ورشة التوعية الأولى: مفاهيم تصنيف البيانات",
        "date": "مايو 2026",
        "count": 16,
        "start_id": 1,
    },
    {
        "title": "ورشة التوعية الثانية: مستويات التصنيف ومتطلباته",
        "date": "مايو 2026",
        "count": 16,
        "start_id": 17,
    },
    {
        "title": "ورشة التوعية الثالثة: آليات التطبيق العملي",
        "date": "يونيو 2026",
        "count": 15,
        "start_id": 33,
    },
]
assert sum(w["count"] for w in WORKSHOPS) == 47


def rtl(p):
    b = OxmlElement("w:bidi")
    b.set(qn("w:val"), "1")
    p._p.get_or_add_pPr().append(b)


def set_rtl_run(run):
    pr = run._element.get_or_add_rPr()
    pr.get_or_add_rFonts().set(qn("w:cs"), AR_FONT)
    e = OxmlElement("w:rtl")
    e.set(qn("w:val"), "1")
    pr.append(e)


def set_table_rtl(table):
    tbl_pr = table._tbl.tblPr
    bv = OxmlElement("w:bidiVisual")
    bv.set(qn("w:val"), "1")
    tbl_pr.append(bv)


def set_cell_background(cell, color_hex):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color_hex)
    cell._tc.get_or_add_tcPr().append(shd)


def add_meta_row(table, label, value):
    row = table.add_row()
    p1 = row.cells[1].paragraphs[0]
    rtl(p1)
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p1.add_run(label)
    r1.font.bold = True
    r1.font.size = Pt(9)
    r1.font.color.rgb = GREY
    set_rtl_run(r1)

    p2 = row.cells[0].paragraphs[0]
    rtl(p2)
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = p2.add_run(value)
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
    set_rtl_run(r2)


def add_paragraph_ar(doc, text, size=9.5, bold=False, italic=False, color=RGBColor(0x26, 0x26, 0x26)):
    p = doc.add_paragraph()
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    set_rtl_run(r)
    return p


def add_heading_ar(doc, text, size=13):
    p = doc.add_paragraph()
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = NAVY
    set_rtl_run(r)
    return p


def make_table(doc, headers, rows, col_widths, center_cols=None):
    center_cols = center_cols or []
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_rtl(table)
    for i, w in enumerate(col_widths):
        table.columns[i].width = w

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        set_cell_background(hdr_cells[i], "1F2A44")
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = hdr_cells[i].paragraphs[0]
        rtl(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_rtl_run(r)

    for idx, row_vals in enumerate(rows):
        row = table.add_row()
        for c, v in enumerate(row_vals):
            cell = row.cells[c]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if idx % 2 == 0:
                set_cell_background(cell, "F2F2F2")
            p = cell.paragraphs[0]
            rtl(p)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c in center_cols else WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(str(v))
            run.font.size = Pt(8.5)
            set_rtl_run(run)
            if c == 0:
                run.font.bold = True
                run.font.color.rgb = NAVY
            else:
                run.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
    return table


doc = docx.Document()
section = doc.sections[0]
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)
section.top_margin = Cm(1.5)
section.bottom_margin = Cm(1.5)

style = doc.styles["Normal"]
style.font.name = AR_FONT
style.font.size = Pt(10)
style.element.rPr.rFonts.set(qn("w:cs"), AR_FONT)

# ============ Cover / banner ============
banner = doc.add_table(rows=1, cols=1)
banner.autofit = True
set_table_rtl(banner)
cell = banner.rows[0].cells[0]
set_cell_background(cell, "1F2A44")
cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
p = cell.paragraphs[0]
rtl(p)
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = p.add_run("سجل حضور ورش التوعية والتدريب الخاصة بتصنيف البيانات")
run.font.size = Pt(18)
run.font.bold = True
run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
set_rtl_run(run)
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(6)

doc.add_paragraph()
add_paragraph_ar(doc, "منصة ذكية لإدارة وقياس الامتثال لمتطلبات المكتب الوطني لإدارة البيانات (NDMO)",
                  size=10.5, bold=True, color=GOLD)
add_heading_ar(doc, "سجل حضور ورش التوعية والتدريب الخاصة بتصنيف البيانات", size=15)
add_paragraph_ar(doc, "دليل داعم — DC.M.2-E04-A", size=11, italic=True, color=GOLD)

doc.add_paragraph()

meta_table = doc.add_table(rows=0, cols=2)
meta_table.autofit = True
set_table_rtl(meta_table)
meta_table.columns[0].width = Cm(11)
meta_table.columns[1].width = Cm(5)
add_meta_row(meta_table, "الجهة", "هيئة الخدمات الحكومية الذكية (SGSA)")
add_meta_row(meta_table, "القسم", "مكتب إدارة البيانات")
add_meta_row(meta_table, "المنصة", "")
add_meta_row(meta_table, "الإطار المرجعي", "إطار حوكمة البيانات الوطني (NDMO)")
add_meta_row(meta_table, "رمز الدليل", "DC.M.2-E04-A")
add_meta_row(meta_table, "المتطلب المرتبط", "DC.M.2 — تنفيذ خطة تصنيف البيانات وخارطة الطريق")
add_meta_row(meta_table, "اسم الدليل", "سجل حضور ورش التوعية والتدريب الخاصة بتصنيف البيانات")
add_meta_row(meta_table, "مالك الوثيقة", "مكتب إدارة البيانات (Data Management Office)")
add_meta_row(meta_table, "الإصدار", "1.0")
add_meta_row(meta_table, "التاريخ", "يونيو 2026")
add_meta_row(meta_table, "حالة الوثيقة", "جاهزة للمراجعة")

doc.add_page_break()

# ============ القسم الأول: بيانات الورشة ============
add_heading_ar(doc, "القسم الأول: بيانات الورشة")
add_paragraph_ar(
    doc,
    "يوثّق هذا الدليل سجل حضور ورش التوعية الثلاث المنفذة ضمن المرحلة الأولى من خطة تصنيف "
    "البيانات، ويُستخدم كمرفق تدقيق داعم لإثبات تنفيذ نشاط التوعية والتدريب ورفع مستوى الوعي "
    "بمتطلبات تصنيف البيانات."
)
doc.add_paragraph()

for w in WORKSHOPS:
    add_paragraph_ar(doc, w["title"], bold=True, size=10.5, color=NAVY)
    add_paragraph_ar(doc, f"التاريخ: {w['date']}", size=9, color=GREY)
    add_paragraph_ar(doc, f"عدد الحضور: {w['count']}", size=9, color=GREY)
    doc.add_paragraph()

doc.add_page_break()

# ============ القسم الثاني: سجل الحضور لكل ورشة ============
add_heading_ar(doc, "القسم الثاني: سجل الحضور لكل ورشة")
doc.add_paragraph()

for w_idx, w in enumerate(WORKSHOPS):
    add_paragraph_ar(doc, w["title"], bold=True, size=11, color=NAVY)
    add_paragraph_ar(doc, f"التاريخ: {w['date']}  —  عدد الحضور: {w['count']}", size=9, color=GREY)
    doc.add_paragraph()

    rows = []
    for i in range(w["count"]):
        participant_id = w["start_id"] + i
        rows.append((
            i + 1,
            NAMES[participant_id - 1],
            dept_for(participant_id - 1),
            "____________",
        ))
    make_table(
        doc,
        ["الرقم", "اسم المشارك", "الإدارة", "التوقيع"],
        rows,
        [Cm(1.8), Cm(4.5), Cm(6.0), Cm(4.0)],
        center_cols=[0, 3],
    )
    if w_idx < len(WORKSHOPS) - 1:
        doc.add_page_break()
    else:
        doc.add_paragraph()

doc.add_page_break()

# ============ القسم الثالث: اعتماد السجل ============
add_heading_ar(doc, "القسم الثالث: اعتماد السجل")
approval_table = doc.add_table(rows=2, cols=4)
approval_table.style = "Table Grid"
set_table_rtl(approval_table)
headers = ["المسمى الوظيفي", "الاسم", "التوقيع", "التاريخ"]
hdr_cells = approval_table.rows[0].cells
for i, h in enumerate(headers):
    set_cell_background(hdr_cells[i], "1F2A44")
    hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = hdr_cells[i].paragraphs[0]
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(h)
    r.font.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    set_rtl_run(r)

vals = ["مدير مكتب إدارة البيانات", "", "", ""]
row_cells = approval_table.rows[1].cells
for i, v in enumerate(vals):
    p = row_cells[i].paragraphs[0]
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(v)
    r.font.size = Pt(9)
    set_rtl_run(r)

doc.add_paragraph()
add_paragraph_ar(
    doc,
    "ملاحظة: تم إعداد هذا السجل ضمن حزمة أدلة تصنيف البيانات الخاصة بهيئة الخدمات الحكومية الذكية (SGSA) التوضيحية وفق نموذج سجلات التدقيق "
    "المعتمد، وتستخدم بيانات نموذجية لغرض عرض آلية توثيق الحضور، ويتم استبدالها بالبيانات "
    "الفعلية عند التطبيق التشغيلي.",
    size=8, italic=True, color=GREY,
)
add_paragraph_ar(
    doc,
    "هذا الدليل مرفق تدقيق مستقل ضمن حزمة أدلة تصنيف البيانات الخاصة بهيئة الخدمات الحكومية الذكية (SGSA) الداعمة لمتطلب DC.M.2، ويستخدم "
    "لإثبات تنفيذ أنشطة التوعية والتدريب الخاصة بتصنيف البيانات.",
    size=8, italic=True, color=GREY,
)

doc.save(OUT_PATH)
print("Saved:", OUT_PATH)
print("Total attendance:", sum(w["count"] for w in WORKSHOPS))
