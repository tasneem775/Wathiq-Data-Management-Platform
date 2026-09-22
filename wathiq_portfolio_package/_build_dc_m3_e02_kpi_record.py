# -*- coding: utf-8 -*-
"""Builds a NEW, standalone supporting evidence record for DC.M.3 (DC-KPI-02) within the
Wathiq/SGSA portfolio context — independent from any existing evidence_repository file.

Purpose: document the data source and calculation method behind the "100%" current value
of DC-KPI-02 (percentage of data-owner assignment) as reported in DC.M.3.docx.

Approved source per the preceding scope-verification audit: DC.M.5.docx ONLY (its scope —
"مجموعات البيانات" only, no records — matches DC-KPI-02's own wording, which never uses
"أصول" or "سجلات"). DC.C.3.1 and the Wathiq-package DC.M.2-E02 register are explicitly
excluded. All dataset names and owner names below are copied verbatim from DC.M.5.docx.
No name is invented. Neither DC.M.3 nor DC.M.5 is modified.
"""
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_PATH = "06_DC_M3_KPI_Evidence/DC.M.3-E02_Data_Owner_Assignment_KPI_Calculation_Record.docx"

NAVY = RGBColor(0x1F, 0x2A, 0x44)
GOLD = RGBColor(0xAB, 0x8D, 0x54)
GREY = RGBColor(0x59, 0x59, 0x59)
GREEN = RGBColor(0x1F, 0x6F, 0x43)
AMBER = RGBColor(0x8A, 0x6D, 0x00)
AR_FONT = "Arial"


def rtl(p):
    b = OxmlElement("w:bidi")
    b.set(qn("w:val"), "1")
    p._p.get_or_add_pPr().append(b)


def set_rtl_run(run):
    pr = run._element.get_or_add_rPr()
    pr.get_or_add_rFonts().set(qn("w:cs"), AR_FONT)
    e = OxmlElement("w:rtl")
    e.set(qn("w:val"), "1")
    pr.append(e)


def set_table_rtl(table):
    tbl_pr = table._tbl.tblPr
    bv = OxmlElement("w:bidiVisual")
    bv.set(qn("w:val"), "1")
    tbl_pr.append(bv)


def set_cell_background(cell, color_hex):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color_hex)
    cell._tc.get_or_add_tcPr().append(shd)


def add_meta_row(table, label, value):
    row = table.add_row()
    p1 = row.cells[1].paragraphs[0]
    rtl(p1)
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p1.add_run(label)
    r1.font.bold = True
    r1.font.size = Pt(9)
    r1.font.color.rgb = GREY
    set_rtl_run(r1)

    p2 = row.cells[0].paragraphs[0]
    rtl(p2)
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = p2.add_run(value)
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
    set_rtl_run(r2)


def add_paragraph_ar(doc, text, size=9.5, bold=False, italic=False, color=RGBColor(0x26, 0x26, 0x26)):
    p = doc.add_paragraph()
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    set_rtl_run(r)
    return p


def add_heading_ar(doc, text, size=13):
    p = doc.add_paragraph()
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = NAVY
    set_rtl_run(r)
    return p


def make_table(doc, headers, rows, col_widths, status_col=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_rtl(table)
    for i, w in enumerate(col_widths):
        table.columns[i].width = w

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        set_cell_background(hdr_cells[i], "1F2A44")
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = hdr_cells[i].paragraphs[0]
        rtl(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_rtl_run(r)

    for idx, row_vals in enumerate(rows):
        row = table.add_row()
        for c, v in enumerate(row_vals):
            cell = row.cells[c]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if idx % 2 == 0:
                set_cell_background(cell, "F2F2F2")
            p = cell.paragraphs[0]
            rtl(p)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c in (status_col or []) else WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(str(v))
            run.font.size = Pt(font_size)
            set_rtl_run(run)
            if c == 0:
                run.font.bold = True
                run.font.color.rgb = NAVY
            elif status_col and c in status_col:
                run.font.bold = True
                run.font.color.rgb = GREEN
            else:
                run.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
    return table


# ============ بيانات مصدرها الحرفي DC.M.5.docx (اسم مجموعة البيانات، مالك البيانات) ============
DATASETS_FROM_DC_M5 = [
    (1, "بيانات الخدمات الحكومية الرقمية", "مدير الخدمات الرقمية"),
    (2, "بيانات التقارير التشغيلية", "مدير التخطيط"),
    (3, "بيانات استطلاعات رأي المستفيدين", "مدير الجودة"),
    (4, "بيانات المستخدمين", "مدير الخدمات الرقمية"),
    (5, "بيانات طلبات المستفيدين", "مدير خدمات المستفيدين"),
    (6, "بيانات البلاغات والشكاوى", "مدير الجودة"),
    (7, "بيانات الموردين", "مدير المشتريات"),
    (8, "بيانات الأصول التقنية", "مدير تقنية المعلومات"),
    (9, "بيانات المراسلات الرسمية", "مدير الشؤون الإدارية"),
    (10, "بيانات الاجتماعات واللجان", "مدير الشؤون الإدارية"),
    (11, "بيانات الموظفين", "مدير الموارد البشرية"),
    (12, "بيانات الرواتب والمزايا", "المدير المالي"),
    (13, "بيانات العقود والاتفاقيات", "مدير المشتريات"),
    (14, "بيانات الصلاحيات وإدارة الهوية", "مسؤول أمن المعلومات"),
    (15, "بيانات المدفوعات المالية", "المدير المالي"),
    (16, "بيانات الميزانية التشغيلية", "المدير المالي"),
    (17, "بيانات التدقيق الداخلي", "مدير التدقيق الداخلي"),
    (18, "بيانات الهوية والتحقق الإلكتروني", "مسؤول أمن المعلومات"),
    (19, "بيانات التكامل مع الجهات الحكومية", "مدير تقنية المعلومات"),
    (20, "بيانات سجلات الأحداث الأمنية", "مسؤول أمن المعلومات"),
]
TOTAL_DATASETS = len(DATASETS_FROM_DC_M5)              # 20 — كما هو في DC.M.5
OWNED_DATASETS = len(DATASETS_FROM_DC_M5)              # 20 — جميعها لديها مالك بيانات محدد في DC.M.5
KPI_VALUE = round(OWNED_DATASETS / TOTAL_DATASETS * 100)  # = 100

doc = docx.Document()
section = doc.sections[0]
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)
section.top_margin = Cm(1.5)
section.bottom_margin = Cm(1.5)

style = doc.styles["Normal"]
style.font.name = AR_FONT
style.font.size = Pt(10)
style.element.rPr.rFonts.set(qn("w:cs"), AR_FONT)

# ============ 1) صفحة الغلاف ============
banner = doc.add_table(rows=1, cols=1)
banner.autofit = True
set_table_rtl(banner)
cell = banner.rows[0].cells[0]
set_cell_background(cell, "1F2A44")
cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
p = cell.paragraphs[0]
rtl(p)
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = p.add_run("سجل احتساب مؤشر نسبة تعيين ملاك البيانات (DC-KPI-02)")
run.font.size = Pt(18)
run.font.bold = True
run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
set_rtl_run(run)
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(6)

doc.add_paragraph()
add_paragraph_ar(doc, "منصة ذكية لإدارة وقياس الامتثال لمتطلبات المكتب الوطني لإدارة البيانات (NDMO)",
                  size=10.5, bold=True, color=GOLD)
add_heading_ar(doc, "سجل احتساب مؤشر نسبة تعيين ملاك البيانات (DC-KPI-02)", size=14.5)
add_paragraph_ar(doc, "Data Owner Assignment KPI Calculation Record", size=11, italic=True, color=GOLD)

doc.add_paragraph()

meta_table = doc.add_table(rows=0, cols=2)
meta_table.autofit = True
set_table_rtl(meta_table)
meta_table.columns[0].width = Cm(11)
meta_table.columns[1].width = Cm(5)
add_meta_row(meta_table, "الجهة", "هيئة الخدمات الحكومية الذكية (SGSA)")
add_meta_row(meta_table, "المنصة", "")
add_meta_row(meta_table, "الإطار المرجعي", "إطار حوكمة البيانات الوطني (NDMO)")
add_meta_row(meta_table, "رمز الدليل", "DC.M.3-E02")
add_meta_row(meta_table, "المتطلب المرتبط", "DC.M.3 — تقرير مراقبة تنفيذ خطة تصنيف البيانات عبر مؤشرات الأداء الرئيسية")
add_meta_row(meta_table, "المؤشر المرتبط", "DC-KPI-02 — نسبة تعيين ملاك البيانات")
add_meta_row(meta_table, "الإصدار", "1.0")
add_meta_row(meta_table, "التاريخ", "يونيو 2026")
add_meta_row(meta_table, "حالة الوثيقة", "جاهزة للاعتماد")

doc.add_page_break()

# ============ 2) المقدمة ============
add_heading_ar(doc, "1.  المقدمة")
add_paragraph_ar(
    doc,
    "يوثّق هذا السجل مصدر البيانات وطريقة احتساب القيمة الحالية لمؤشر الأداء الرئيسي "
    "DC-KPI-02 (نسبة تعيين ملاك البيانات)، الوارد في بطاقة المؤشر ضمن تقرير "
    "DC.M.3 — تقرير مراقبة تنفيذ خطة تصنيف البيانات عبر مؤشرات الأداء الرئيسية، "
    "والذي أظهر قيمة حالية قدرها 100% مقابل هدف 100%."
)
add_paragraph_ar(
    doc,
    "يعتمد هذا السجل حصراً على البيانات الفعلية الواردة في سجل DC.M.5 — القائمة الحالية "
    "لمجموعات البيانات المصنفة. لا يستخدم هذا السجل تقرير الجرد DC.C.3.1 ولا سجل ملاك "
    "البيانات ضمن حزمة أدلة تصنيف البيانات الخاصة بهيئة الخدمات الحكومية الذكية (SGSA) للمتطلب DC.M.2، إذ إن نطاق مؤشر DC-KPI-02 كما ورد "
    "حرفياً في بطاقته ضمن DC.M.3 مقتصر على «مجموعات البيانات» دون أي ذكر لكلمة «أصول» "
    "أو «سجلات»، وهو ما يطابق نطاق DC.M.5 حصراً (مجموعات بيانات فقط)، بخلاف DC.C.3.1 "
    "الذي يغطي نطاقاً أوسع (مجموعات بيانات وسجلات معاً)."
)

doc.add_paragraph()

# ============ 3) بيانات الاحتساب ============
add_heading_ar(doc, "2.  بيانات الاحتساب")
meta2 = doc.add_table(rows=0, cols=2)
meta2.autofit = True
set_table_rtl(meta2)
meta2.columns[0].width = Cm(11)
meta2.columns[1].width = Cm(5)
add_meta_row(meta2, "مصدر البيانات المستخدم", "DC.M.5 — القائمة الحالية لمجموعات البيانات المصنفة (الإصدار 1.0، يونيو 2026)")
add_meta_row(meta2, "فترة القياس", "الربع الثاني 2026 (أبريل – يونيو 2026)")
add_meta_row(
    meta2, "نطاق القياس",
    "جميع مجموعات البيانات الواردة في سجل DC.M.5 حتى تاريخ إصداره (20 مجموعة بيانات)، "
    "التزاماً بنطاق مؤشر DC-KPI-02 كما ورد حرفياً في بطاقته («المجموعات» فقط)."
)
add_meta_row(meta2, "تاريخ استخراج البيانات", "يونيو 2026 (تاريخ إصدار DC.M.5)")

doc.add_paragraph()

# ============ 4) جدول مجموعات البيانات وملاكها ============
add_heading_ar(doc, "3.  مجموعات البيانات وملاكها (وفق سجل DC.M.5)")
add_paragraph_ar(doc, "يعرض الجدول التالي جميع مجموعات البيانات الـ20 الواردة في سجل DC.M.5 مع مالك البيانات المحدد لكل منها:")
rows = [
    (f"{n:02d}", name, owner)
    for n, name, owner in DATASETS_FROM_DC_M5
]
make_table(
    doc,
    ["الرقم التسلسلي", "اسم مجموعة البيانات", "مالك البيانات"],
    rows,
    [Cm(2.6), Cm(7.0), Cm(6.4)],
    font_size=8.5,
)

doc.add_paragraph()

# ============ 5) حساب المؤشر ============
add_heading_ar(doc, "4.  حساب المؤشر")
make_table(
    doc,
    ["البند", "القيمة"],
    [
        ("إجمالي مجموعات البيانات (وفق سجل DC.M.5)", str(TOTAL_DATASETS)),
        ("عدد المجموعات ذات المالك المحدد", str(OWNED_DATASETS)),
        ("معادلة المؤشر", "عدد المجموعات ذات المالك المحدد ÷ إجمالي المجموعات × 100"),
        ("تطبيق المعادلة", f"{OWNED_DATASETS} ÷ {TOTAL_DATASETS} × 100"),
        ("القيمة النهائية للمؤشر", f"{KPI_VALUE}%"),
    ],
    [Cm(9.0), Cm(7.0)],
    status_col=[1],
)
doc.add_paragraph()
add_paragraph_ar(
    doc,
    f"القيمة الناتجة ({KPI_VALUE}%) مطابقة للقيمة الحالية الواردة في بطاقة مؤشر DC-KPI-02 "
    "ضمن تقرير DC.M.3."
)

doc.add_paragraph()

# ============ 6) التحقق والاعتماد ============
add_heading_ar(doc, "5.  التحقق والاعتماد")
add_paragraph_ar(
    doc,
    "تمت مراجعة هذا السجل من مكتب إدارة البيانات، وتم التحقق من مطابقة أسماء مجموعات "
    "البيانات وملاكها الواردة فيه لمحتوى سجل DC.M.5 دون أي إضافة أو حذف أو تعديل."
)
doc.add_paragraph()

approval_table = doc.add_table(rows=2, cols=4)
approval_table.style = "Table Grid"
set_table_rtl(approval_table)
headers = ["صاحب الصلاحية", "المسمى الوظيفي", "التوقيع", "التاريخ"]
hdr_cells = approval_table.rows[0].cells
for i, h in enumerate(headers):
    set_cell_background(hdr_cells[i], "1F2A44")
    hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = hdr_cells[i].paragraphs[0]
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(h)
    r.font.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    set_rtl_run(r)

vals = ["________________", "مدير مكتب إدارة البيانات", "________________", "________________"]
row_cells = approval_table.rows[1].cells
for i, v in enumerate(vals):
    p = row_cells[i].paragraphs[0]
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(v)
    r.font.size = Pt(9)
    set_rtl_run(r)

doc.add_paragraph()
add_paragraph_ar(
    doc,
    "ملاحظة: هذا السجل دليل داعم مستقل ضمن حزمة أدلة تصنيف البيانات الخاصة بهيئة الخدمات الحكومية الذكية (SGSA)، أُعدّ لتوثيق مصدر بيانات "
    "وطريقة احتساب مؤشر DC-KPI-02 دون تعديل أي محتوى في تقرير DC.M.3 أو سجل DC.M.5.",
    size=8, italic=True, color=GREY,
)

doc.save(OUT_PATH)
print("Saved:", OUT_PATH)
print("TOTAL_DATASETS:", TOTAL_DATASETS)
print("OWNED_DATASETS:", OWNED_DATASETS)
print("KPI_VALUE:", KPI_VALUE)
