# -*- coding: utf-8 -*-
"""Builds the supporting Excel evidence sheet for DC.M.3-E03
(DC-KPI-03 — نسبة إكمال التوعية بتصنيف البيانات).

This is an Evidence Record, not a Calculation Record: it documents the value (90%) as it
appears verbatim in the real evidence_repository file DC.M.2.docx. No employee-target total
is invented, and no independent recalculation of the KPI is claimed.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT_PATH = "06_DC_M3_KPI_Evidence/DC.M.3-E03_Training_Awareness_KPI_Evidence_Record.xlsx"

NAVY = "1F2A44"
GOLD = "AB8654"
LIGHT_GREY = "F2F2F2"
WHITE = "FFFFFF"
GREEN = "1F6F43"
AR_FONT = "Arial"

wb = Workbook()
ws = wb.active
ws.title = "إثبات DC-KPI-03"
ws.sheet_view.showGridLines = False
ws.sheet_view.rightToLeft = True

thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

N_COLS = 5
r = 1


def merged_text(row, text, size, bold, color, fill=None, height=None, italic=False, span=N_COLS):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name=AR_FONT, size=size, bold=bold, italic=italic, color=color)
    cell.alignment = Alignment(horizontal="right", vertical="center", indent=1, wrap_text=True)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if height:
        ws.row_dimensions[row].height = height
    return row + 1


def section_heading(row, text, span=N_COLS):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name=AR_FONT, size=12.5, bold=True, color=NAVY)
    cell.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    cell.border = Border(bottom=Side(style="medium", color=GOLD))
    ws.row_dimensions[row].height = 22
    return row + 1


def spacer(row, h=10):
    ws.row_dimensions[row].height = h
    return row + 1


def meta_row(row, label, value):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    c1 = ws.cell(row=row, column=1, value=label)
    c1.font = Font(name=AR_FONT, size=10, bold=True, color="595959")
    c1.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=N_COLS)
    c2 = ws.cell(row=row, column=3, value=value)
    c2.font = Font(name=AR_FONT, size=10, color="262626")
    c2.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    return row + 1


# ============ Title banner ============
r = merged_text(r, "سجل إثبات مؤشر إكمال التوعية بتصنيف البيانات (DC-KPI-03)", 18, True, WHITE, fill=NAVY, height=32)
r = merged_text(r, "منصة ذكية لإدارة وقياس الامتثال لمتطلبات المكتب الوطني لإدارة البيانات (NDMO)", 10.5, True, WHITE, fill=NAVY, height=20)
r = spacer(r, 6)
r = merged_text(r, "سجل إثبات مؤشر إكمال التوعية بتصنيف البيانات (DC-KPI-03)", 14, True, NAVY, height=24)
r = merged_text(r, "دليل داعم — DC.M.3-E03 — Training Awareness KPI Evidence Record", 11, True, GOLD, italic=True, height=18)
r = spacer(r, 10)

# ============ بيانات الوثيقة ============
r = section_heading(r, "بيانات الوثيقة")
r = spacer(r, 4)
r = meta_row(r, "الجهة", "هيئة الخدمات الحكومية الذكية (SGSA)")
r = meta_row(r, "المنصة", "")
r = meta_row(r, "الإطار المرجعي", "إطار حوكمة البيانات الوطني (NDMO)")
r = meta_row(r, "رمز الدليل", "DC.M.3-E03")
r = meta_row(r, "المتطلب المرتبط", "DC.M.3 — تقرير مراقبة تنفيذ خطة تصنيف البيانات عبر مؤشرات الأداء الرئيسية")
r = meta_row(r, "المؤشر المرتبط", "DC-KPI-03 — نسبة إكمال التوعية بتصنيف البيانات")
r = meta_row(r, "الإصدار", "1.0")
r = meta_row(r, "التاريخ", "يونيو 2026")
r = meta_row(r, "حالة الوثيقة", "جاهزة للاعتماد")
r = spacer(r, 16)

# ============ جدول بيانات الإثبات ============
r = section_heading(r, "بيانات الإثبات (مصدرها DC.M.2)")
r = spacer(r, 4)

headers = ["المؤشر", "القيمة", "مصدر القيمة", "فترة القياس", "ملاحظات التحقق"]
header_row = r
for c, h in enumerate(headers, start=1):
    cell = ws.cell(row=header_row, column=c, value=h)
    cell.font = Font(name=AR_FONT, size=10, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border
ws.row_dimensions[header_row].height = 30
r += 1

data_row = [
    "DC-KPI-03 — نسبة إكمال التوعية بتصنيف البيانات",
    "90%",
    "DC.M.2 — تقرير حالة تنفيذ خطة تصنيف البيانات",
    "حتى يونيو 2026",
    "قيمة واردة/مذكورة في DC.M.2؛ لا يمثل هذا السجل إعادة احتساب مستقل للمؤشر (لا يتوفر رقم إجمالي المستهدفين في المصدر)",
]
for c, v in enumerate(data_row, start=1):
    cell = ws.cell(row=r, column=c, value=v)
    cell.border = border
    cell.fill = PatternFill("solid", fgColor=LIGHT_GREY)
    center_cols = (2,)
    cell.alignment = Alignment(
        horizontal="center" if c in center_cols else "right",
        vertical="center", wrap_text=True,
        indent=0 if c in center_cols else 1,
    )
    if c == 1:
        cell.font = Font(name=AR_FONT, size=9.5, bold=True, color=NAVY)
    elif c == 2:
        cell.font = Font(name=AR_FONT, size=11, bold=True, color=GREEN)
    else:
        cell.font = Font(name=AR_FONT, size=9, color="262626")
ws.row_dimensions[r].height = 60
r += 1

r = spacer(r, 16)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=N_COLS)
note1 = ws.cell(
    row=r, column=1,
    value="هذا الملف يوثق بيانات الإثبات الواردة في DC.M.2 ولا يمثل إعادة احتساب مستقل للمؤشر.",
)
note1.font = Font(name=AR_FONT, size=9, italic=True, color="7F7F7F")
note1.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True, indent=1)
ws.row_dimensions[r].height = 22
r += 1

r = spacer(r, 14)

# ============ الاعتماد ============
r = section_heading(r, "التحقق والاعتماد")
r = spacer(r, 4)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=N_COLS)
review_cell = ws.cell(
    row=r, column=1,
    value="تمت مراجعة هذا السجل والتحقق من مطابقة البيانات الواردة فيه لمحتوى DC.M.2 دون إضافة أو حذف أو تعديل.",
)
review_cell.font = Font(name=AR_FONT, size=10, color="262626")
review_cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True, indent=1)
ws.row_dimensions[r].height = 26
r += 1
r = spacer(r, 6)

approval_headers = ["صاحب الصلاحية", "المسمى الوظيفي", "التوقيع", "التاريخ"]
col_spans = [(1, 1), (2, 2), (3, 4), (5, 5)]
approval_header_row = r
for (sc, ec), h in zip(col_spans, approval_headers):
    ws.merge_cells(start_row=approval_header_row, start_column=sc, end_row=approval_header_row, end_column=ec)
    cell = ws.cell(row=approval_header_row, column=sc, value=h)
    cell.font = Font(name=AR_FONT, size=10, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    for cc in range(sc, ec + 1):
        ws.cell(row=approval_header_row, column=cc).border = border
ws.row_dimensions[approval_header_row].height = 20
r += 1

approval_values = ["____________", "مدير مكتب إدارة البيانات", "____________", "____________"]
for (sc, ec), v in zip(col_spans, approval_values):
    ws.merge_cells(start_row=r, start_column=sc, end_row=r, end_column=ec)
    cell = ws.cell(row=r, column=sc, value=v)
    cell.font = Font(name=AR_FONT, size=9.5, color="262626")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    for cc in range(sc, ec + 1):
        ws.cell(row=r, column=cc).border = border
ws.row_dimensions[r].height = 26
r += 1

r = spacer(r, 14)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=N_COLS)
note2 = ws.cell(
    row=r, column=1,
    value="ملاحظة: هذا الملف دليل داعم مستقل ضمن حزمة أدلة تصنيف البيانات الخاصة بهيئة الخدمات الحكومية الذكية (SGSA)، أُعدّ لتوثيق مصدر وقابلية تتبع قيمة مؤشر DC-KPI-03 دون تعديل أي محتوى في تقرير DC.M.3 أو تقرير DC.M.2.",
)
note2.font = Font(name=AR_FONT, size=8, italic=True, color="7F7F7F")
note2.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True, indent=1)
ws.row_dimensions[r].height = 26

# ============ Column widths ============
widths = [34, 10, 34, 16, 44]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

wb.save(OUT_PATH)
print("Saved:", OUT_PATH)
