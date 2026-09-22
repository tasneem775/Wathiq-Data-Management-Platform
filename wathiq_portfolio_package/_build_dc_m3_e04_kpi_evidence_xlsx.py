# -*- coding: utf-8 -*-
"""Builds the supporting Excel evidence sheet for DC.M.3-E04
(DC-KPI-04 — نسبة الالتزام بإجراءات تصنيف البيانات).

Simple KPI Evidence sheet as specified: documents baseline/current value as stated
verbatim in the real evidence_repository file DC.M.4.docx. No case counts are invented,
and no independent recalculation of the KPI is claimed.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT_PATH = "06_DC_M3_KPI_Evidence/DC.M.3-E04_Compliance_Procedures_KPI_Evidence_Record.xlsx"

NAVY = "1F2A44"
WHITE = "FFFFFF"
GREEN = "1F6F43"
LIGHT_GREY = "F2F2F2"
AR_FONT = "Arial"

wb = Workbook()
ws = wb.active
ws.title = "KPI Evidence"

thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

headers = [
    "Indicator Code",
    "Indicator Name",
    "Baseline",
    "Current Value",
    "Source Document",
    "Recalculation Status",
]
for c, h in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = Font(name=AR_FONT, size=10.5, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border
ws.row_dimensions[1].height = 24

row_vals = [
    "DC-KPI-04",
    "نسبة الالتزام بإجراءات تصنيف البيانات",
    "65%",
    "88%",
    "DC.M.4",
    "Not recalculable from available evidence",
]
for c, v in enumerate(row_vals, start=1):
    cell = ws.cell(row=2, column=c, value=v)
    cell.border = border
    cell.fill = PatternFill("solid", fgColor=LIGHT_GREY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if c == 1:
        cell.font = Font(name=AR_FONT, size=10, bold=True, color=NAVY)
    elif c in (3, 4):
        cell.font = Font(name=AR_FONT, size=11, bold=True, color=GREEN)
    else:
        cell.font = Font(name=AR_FONT, size=10, color="262626")
ws.row_dimensions[2].height = 30

widths = [14, 40, 12, 14, 16, 34]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

wb.save(OUT_PATH)
print("Saved:", OUT_PATH)
