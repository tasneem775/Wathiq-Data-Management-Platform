"""Builds the Evidence Mapping Inventory (Excel) for the Wathiq governance portfolio package.

Synthetic case-study evidence only — not a real government document.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT_PATH = "03_Evidence_Mapping_Inventory/Evidence_Mapping_Inventory.xlsx"

NAVY = "1F2A44"
LIGHT_GREY = "F2F2F2"
WHITE = "FFFFFF"
AMBER = "8A6D00"
GREEN = "1F6F43"

ROWS = [
    (
        "SUP-EV-001",
        "DC.M.2",
        "Primary Narrative Evidence (Reference Only)",
        "DC.M.2 — Implementation Status Report (primary evidence; outside the scope of this supporting package)",
        "Referenced — No Additional Artifact Required",
    ),
    (
        "SUP-EV-002",
        "DC.M.3",
        "KPI Supporting Data",
        "KPI_Supporting_Data_Report.xlsx / .pdf",
        "Completed",
    ),
    (
        "SUP-EV-003",
        "DC.M.4",
        "Improvement Evidence",
        "Data_Classification_Improvement_Evidence.docx / .pdf",
        "Completed",
    ),
    (
        "SUP-EV-004",
        "DC.C.4.1",
        "Classification Metrics (KPI Supporting Data)",
        "KPI_Supporting_Data_Report.xlsx / .pdf",
        "Completed",
    ),
    (
        "SUP-EV-005",
        "DC.M.12",
        "KPI Monitoring Evidence",
        "KPI_Supporting_Data_Report.xlsx / .pdf",
        "Completed",
    ),
]

wb = Workbook()
ws = wb.active
ws.title = "Evidence Mapping Inventory"
ws.sheet_view.showGridLines = False

thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

ws.merge_cells("A1:E1")
ws["A1"] = "سجل ربط الأدلة — حوكمة تصنيف البيانات"
ws["A1"].font = Font(name="Calibri", size=16, bold=True, color=WHITE)
ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws.row_dimensions[1].height = 30

ws.merge_cells("A2:E2")
ws["A2"] = "Evidence Mapping Inventory — Data Classification Governance"
ws["A2"].font = Font(name="Calibri", size=12, bold=True, color=NAVY)
ws["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws.row_dimensions[2].height = 20

meta_rows = [
    ("Entity (Synthetic Case Study)", "Smart Government Services Authority (SGSA)"),
    ("Department", "Data Management Office"),
    ("Framework", "NDMO Data Governance Framework"),
    ("Purpose", "Links supporting evidence artifacts to their corresponding requirements."),
]
r = 4
for label, value in meta_rows:
    ws.cell(row=r, column=1, value=label).font = Font(name="Calibri", size=10, bold=True, color="595959")
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    ws.cell(row=r, column=2, value=value).font = Font(name="Calibri", size=10, color="262626")
    r += 1

r += 1
table_header_row = r
headers = ["Evidence ID", "Requirement", "Evidence Type", "Artifact Name", "Status"]
for c, h in enumerate(headers, start=1):
    cell = ws.cell(row=table_header_row, column=c, value=h)
    cell.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border
ws.row_dimensions[table_header_row].height = 22

r += 1
for i, (eid, req, etype, artifact, status) in enumerate(ROWS):
    row_fill = LIGHT_GREY if i % 2 == 0 else WHITE
    values = [eid, req, etype, artifact, status]
    for c, v in enumerate(values, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.border = border
        cell.alignment = Alignment(
            horizontal="center" if c in (1, 2, 5) else "left", vertical="center", wrap_text=True
        )
        cell.fill = PatternFill("solid", fgColor=row_fill)
        if c == 2:
            cell.font = Font(name="Calibri", size=10, bold=True, color=NAVY)
        elif c == 5:
            color = AMBER if "Referenced" in v else GREEN
            cell.font = Font(name="Calibri", size=10, bold=True, color=color)
        else:
            cell.font = Font(name="Calibri", size=10, color="262626")
    ws.row_dimensions[r].height = 32
    r += 1

r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
note = ws.cell(
    row=r,
    column=1,
    value=(
        "Note: This document is synthetic case-study evidence produced for the Wathiq governance "
        "portfolio demonstration. It does not represent an actual government entity, a real "
        "compliance submission, or verified production data."
    ),
)
note.font = Font(name="Calibri", size=8, italic=True, color="7F7F7F")
note.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
ws.row_dimensions[r].height = 28

widths = [14, 14, 30, 48, 30]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

wb.save(OUT_PATH)
print("Saved:", OUT_PATH)
